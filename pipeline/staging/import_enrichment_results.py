"""
import_enrichment_results.py
Imports verified Gemini enrichment results from XLSX into data/staging.db.

STRICT CONTRACT:
1. Does NOT call Gemini API.
2. Does NOT modify the source results XLSX.
3. Does NOT alter question text, options, gaps, or explanations.
4. Preserves existing 182-question Gemini checkpoint answers untouched.
5. Populates correct_answer, accepted_answers, and is_correct for the remaining 5,614 questions.
6. All updates execute in a single atomic database transaction (full rollback on error).
7. Strict validation: single_choice must have exactly 1 matching option; multiple_choice >= 1.
"""

from pathlib import Path
import argparse
import json
import logging
import sqlite3
import sys
from typing import Dict, Any, List, Set, Tuple
import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
DEFAULT_RESULTS = Path.home() / "Desktop" / "english_cms_gemini_remaining_5614 (3).xlsx"
FALLBACK_RESULTS = REPO_ROOT / "data" / "gemini" / "english_cms_gemini_remaining_5614_evaluated.xlsx"
DEFAULT_DB = REPO_ROOT / "data" / "staging.db"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s"
)
logger = logging.getLogger("import_enrichment_results")

# Known prompt-anomaly overrides
KNOWN_OVERRIDES: Dict[str, Dict[str, Any]] = {
    # QID 12660: quiz-1472 sentence 9 had missing word box in prompt; target past tense from source box 'always pick' -> 'were always picking'
    "12660": {"gap_1": "were always picking"},
    # QID 187: quiz-25 sentence 4 had 'as soon as' in example, but select options list had 'when'
    "187": {"gap_1": "when"}
}


def load_enrichment_results(xlsx_path: Path) -> List[Dict[str, Any]]:
    """Loads and validates the 5,614 rows from the evaluated Gemini workbook."""
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Enrichment results XLSX not found at: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Gemini_Enrichment" not in wb.sheetnames:
        raise ValueError(f"Missing 'Gemini_Enrichment' sheet in {xlsx_path}. Found: {wb.sheetnames}")

    ws = wb["Gemini_Enrichment"]

    records = []
    seen_qids: Set[str] = set()

    for idx, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if r[0] is None:
            continue
        qid = str(int(r[0]))
        if qid in seen_qids:
            raise ValueError(f"Duplicate question_id {qid} found at row {idx}")
        seen_qids.add(qid)

        exercise_id = str(r[1] or "")
        lesson_id = str(r[2] or "")
        order = int(r[3] or 0)
        response_model = str(r[4] or "").strip()
        instruction = str(r[5] or "")
        sentence = str(r[6] or "")
        options_context = str(r[7] or "")

        # Column J (idx 9) is raw evaluated value; fallback to col I (idx 8)
        val = r[9] if (len(r) > 9 and r[9] is not None) else (r[8] if len(r) > 8 else None)
        if val is None and qid not in KNOWN_OVERRIDES:
            raise ValueError(f"Row {idx} (QID {qid}) has null answer value")

        if qid in KNOWN_OVERRIDES:
            parsed_answer = KNOWN_OVERRIDES[qid]
        else:
            val_str = str(val).strip()
            try:
                parsed_answer = json.loads(val_str)
            except json.JSONDecodeError as exc:
                raise ValueError(f"Row {idx} (QID {qid}) invalid JSON: {val_str}") from exc

        records.append({
            "row_idx": idx,
            "question_id": qid,
            "exercise_id": exercise_id,
            "lesson_id": lesson_id,
            "order": order,
            "response_model": response_model,
            "instruction": instruction,
            "sentence": sentence,
            "options_context": options_context,
            "parsed_answer": parsed_answer,
        })

    return records


def match_option(target: str, options: List[Tuple[str, str, str]]) -> List[str]:
    """Matches target string against list of (option_id, text, value).
    Returns list of matching option_ids (checks exact, then case-insensitive)."""
    t_clean = target.strip()
    t_lower = t_clean.lower()

    # 1. Exact match on text or value
    exact = [opt_id for opt_id, text, val in options if (text or "").strip() == t_clean or (val or "").strip() == t_clean]
    if exact:
        return exact

    # 2. Case-insensitive match on text or value
    case_ins = [opt_id for opt_id, text, val in options if (text or "").strip().lower() == t_lower or (val or "").strip().lower() == t_lower]
    return case_ins


def import_enrichment_results(
    results_path: Path,
    db_path: Path,
    dry_run: bool = False
) -> Dict[str, Any]:
    """Imports enrichment answers into SQLite staging database within an atomic transaction."""
    records = load_enrichment_results(results_path)
    logger.info(f"Loaded {len(records)} records from enrichment workbook.")
    if len(records) != 5614:
        logger.warning(f"Expected 5,614 records, found {len(records)}")

    if not db_path.exists():
        raise FileNotFoundError(f"Staging database not found at: {db_path}")

    conn = sqlite3.connect(str(db_path))
    conn.execute("PRAGMA foreign_keys = ON;")
    cursor = conn.cursor()

    report: Dict[str, Any] = {
        "workbook_questions": len(records),
        "total_questions_in_staging": 0,
        "matched_questions": 0,
        "unmatched_questions": 0,
        "single_choice_questions_answered": 0,
        "multiple_choice_questions_answered": 0,
        "gap_questions_answered": 0,
        "total_gaps_answered_in_batch": 0,
        "total_options_updated_in_batch": 0,
        "total_answered_questions": 0,
        "remaining_unanswered_questions": 0,
        "remaining_unanswered_gaps": 0,
        "remaining_unanswered_options": 0,
        "validation_errors": []
    }

    try:
        cursor.execute("BEGIN TRANSACTION;")

        # 1. Fetch initial checkpoint status and ensure 0 collisions
        cursor.execute("""
            SELECT DISTINCT question_id FROM staging_gaps WHERE correct_answer IS NOT NULL
            UNION
            SELECT DISTINCT question_id FROM staging_options WHERE is_correct IS NOT NULL
        """)
        checkpoint_qids = set(row[0] for row in cursor.fetchall())
        logger.info(f"Existing answered checkpoint questions in staging: {len(checkpoint_qids)}")
        if len(checkpoint_qids) != 182:
            logger.warning(f"Expected 182 existing answered questions, found {len(checkpoint_qids)}")

        # Snapshot checkpoint values to ensure zero modification
        cursor.execute("""
            SELECT gap_id, correct_answer, accepted_answers FROM staging_gaps
            WHERE question_id IN (SELECT value FROM json_each(?))
        """, (json.dumps(list(checkpoint_qids)),))
        checkpoint_gaps_snapshot = cursor.fetchall()

        cursor.execute("""
            SELECT option_id, is_correct FROM staging_options
            WHERE question_id IN (SELECT value FROM json_each(?))
        """, (json.dumps(list(checkpoint_qids)),))
        checkpoint_opts_snapshot = cursor.fetchall()

        new_qids: Set[str] = set()

        for item in records:
            qid = item["question_id"]
            model = item["response_model"]
            parsed_ans = item["parsed_answer"]
            new_qids.add(qid)

            if qid in checkpoint_qids:
                report["validation_errors"].append(f"QID {qid} collides with existing 182 checkpoint questions")
                continue

            # Check question exists in staging
            cursor.execute(
                "SELECT question_id, response_model FROM staging_questions WHERE question_id = ?;",
                (qid,)
            )
            q_row = cursor.fetchone()
            if not q_row:
                err = f"Question ID {qid} not found in staging_questions"
                report["validation_errors"].append(err)
                report["unmatched_questions"] += 1
                continue

            report["matched_questions"] += 1

            # Process by model
            if model == "gap":
                report["gap_questions_answered"] += 1
                cursor.execute(
                    "SELECT gap_id, gap_order, input_control FROM staging_gaps WHERE question_id = ? ORDER BY gap_order;",
                    (qid,)
                )
                db_gaps = cursor.fetchall()
                if not db_gaps:
                    err = f"No gaps found in staging for gap question {qid}"
                    report["validation_errors"].append(err)
                    continue

                for g_key, g_ans in parsed_ans.items():
                    try:
                        g_order = int(g_key.replace("gap_", "")) if "gap_" in g_key else int(g_key)
                    except ValueError:
                        g_order = 1

                    matched_gap = None
                    for dg in db_gaps:
                        if dg[1] == g_order or dg[0] == g_key:
                            matched_gap = dg
                            break

                    if not matched_gap:
                        err = f"Gap key '{g_key}' for QID {qid} could not be matched to staging gaps"
                        report["validation_errors"].append(err)
                        continue

                    gap_id, gap_order, input_control = matched_gap
                    accepted_json = json.dumps([str(g_ans)])

                    cursor.execute("""
                        UPDATE staging_gaps
                        SET correct_answer = ?, accepted_answers = ?
                        WHERE gap_id = ?;
                    """, (str(g_ans), accepted_json, gap_id))
                    report["total_gaps_answered_in_batch"] += 1

                    if input_control == "select":
                        cursor.execute(
                            "SELECT option_id, text, value FROM staging_options WHERE gap_id = ?;",
                            (gap_id,)
                        )
                        gap_opts = cursor.fetchall()
                        matched_opt_ids = match_option(str(g_ans), gap_opts)

                        if len(matched_opt_ids) != 1:
                            err = f"QID {qid} Gap {gap_id}: Expected exactly 1 match for '{g_ans}', found {len(matched_opt_ids)} among {[o[1] for o in gap_opts]}"
                            report["validation_errors"].append(err)

                        for opt_id, _, _ in gap_opts:
                            is_corr = 1 if opt_id in matched_opt_ids else 0
                            cursor.execute("""
                                UPDATE staging_options
                                SET is_correct = ?
                                WHERE option_id = ?;
                            """, (is_corr, opt_id))
                            report["total_options_updated_in_batch"] += 1

            elif model in ("single_choice", "multiple_choice"):
                if model == "single_choice":
                    report["single_choice_questions_answered"] += 1
                else:
                    report["multiple_choice_questions_answered"] += 1

                if isinstance(parsed_ans, dict):
                    if "answers" in parsed_ans:
                        ans_val = parsed_ans["answers"]
                    elif "answer" in parsed_ans:
                        ans_val = parsed_ans["answer"]
                    elif "selected" in parsed_ans:
                        ans_val = parsed_ans["selected"]
                    else:
                        ans_val = list(parsed_ans.values())[0]
                else:
                    ans_val = parsed_ans

                if isinstance(ans_val, list):
                    ans_list = [str(x) for x in ans_val]
                else:
                    ans_list = [str(ans_val)]

                cursor.execute(
                    "SELECT option_id, text, value FROM staging_options WHERE question_id = ?;",
                    (qid,)
                )
                q_opts = cursor.fetchall()
                if not q_opts:
                    err = f"No options found in staging for choice question {qid}"
                    report["validation_errors"].append(err)
                    continue

                matched_opt_ids = set()
                for target_item in ans_list:
                    matched_for_item = match_option(target_item, q_opts)
                    matched_opt_ids.update(matched_for_item)

                if model == "single_choice" and len(matched_opt_ids) != 1:
                    err = f"QID {qid} (single_choice): Expected exactly 1 matching option, found {len(matched_opt_ids)} (targets: {ans_list}, options: {[o[1] for o in q_opts]})"
                    report["validation_errors"].append(err)
                elif model == "multiple_choice" and len(matched_opt_ids) == 0:
                    err = f"QID {qid} (multiple_choice): Expected at least 1 matching option, found 0 (targets: {ans_list}, options: {[o[1] for o in q_opts]})"
                    report["validation_errors"].append(err)

                for opt_id, _, _ in q_opts:
                    is_corr = 1 if opt_id in matched_opt_ids else 0
                    cursor.execute("""
                        UPDATE staging_options
                        SET is_correct = ?
                        WHERE option_id = ?;
                    """, (is_corr, opt_id))
                    report["total_options_updated_in_batch"] += 1

        # Boundary checks: Checkpoint 182 questions must remain byte-for-byte identical
        cursor.execute("""
            SELECT gap_id, correct_answer, accepted_answers FROM staging_gaps
            WHERE question_id IN (SELECT value FROM json_each(?))
        """, (json.dumps(list(checkpoint_qids)),))
        current_checkpoint_gaps = cursor.fetchall()
        if current_checkpoint_gaps != checkpoint_gaps_snapshot:
            report["validation_errors"].append("CRITICAL: Existing 182-question checkpoint gaps were altered!")

        cursor.execute("""
            SELECT option_id, is_correct FROM staging_options
            WHERE question_id IN (SELECT value FROM json_each(?))
        """, (json.dumps(list(checkpoint_qids)),))
        current_checkpoint_opts = cursor.fetchall()
        if current_checkpoint_opts != checkpoint_opts_snapshot:
            report["validation_errors"].append("CRITICAL: Existing 182-question checkpoint options were altered!")

        # Global database metrics
        cursor.execute("SELECT COUNT(*) FROM staging_questions;")
        report["total_questions_in_staging"] = cursor.fetchone()[0]

        cursor.execute("SELECT COUNT(*) FROM staging_gaps WHERE correct_answer IS NULL;")
        report["remaining_unanswered_gaps"] = cursor.fetchone()[0]

        cursor.execute("SELECT COUNT(*) FROM staging_options WHERE is_correct IS NULL;")
        report["remaining_unanswered_options"] = cursor.fetchone()[0]

        cursor.execute("""
            SELECT COUNT(DISTINCT question_id) FROM (
                SELECT question_id FROM staging_gaps WHERE correct_answer IS NOT NULL
                UNION
                SELECT question_id FROM staging_options WHERE is_correct IS NOT NULL
            )
        """)
        report["total_answered_questions"] = cursor.fetchone()[0]
        report["remaining_unanswered_questions"] = report["total_questions_in_staging"] - report["total_answered_questions"]

        if report["validation_errors"]:
            logger.error(f"Validation failed with {len(report['validation_errors'])} errors. Rolling back transaction.")
            cursor.execute("ROLLBACK;")
            raise RuntimeError(f"Import aborted due to validation errors: {report['validation_errors'][:5]}")

        if dry_run:
            logger.info("Dry run requested — rolling back transaction.")
            cursor.execute("ROLLBACK;")
        else:
            conn.commit()
            logger.info("Enrichment results import committed successfully.")

    except Exception as e:
        conn.rollback()
        logger.error(f"Transaction rolled back: {e}")
        raise
    finally:
        conn.close()

    return report


def main():
    parser = argparse.ArgumentParser(description="Import 5,614 Gemini enrichment results into staging database.")
    parser.add_argument("--results", type=Path, default=DEFAULT_RESULTS, help="Path to enrichment results XLSX")
    parser.add_argument("--db", type=Path, default=DEFAULT_DB, help="Path to SQLite staging DB")
    parser.add_argument("--dry-run", action="store_true", help="Validate without committing changes")
    args = parser.parse_args()

    results_file = args.results
    if not results_file.exists() and FALLBACK_RESULTS.exists():
        results_file = FALLBACK_RESULTS

    report = import_enrichment_results(results_file, args.db, dry_run=args.dry_run)

    print("\n" + "=" * 60)
    print("GEMINI ENRICHMENT IMPORT REPORT")
    print("=" * 60)
    for k, v in report.items():
        if k != "validation_errors":
            print(f"  {k}: {v}")
    print(f"  validation_errors_count: {len(report['validation_errors'])}")
    print("=" * 60)


if __name__ == "__main__":
    main()
