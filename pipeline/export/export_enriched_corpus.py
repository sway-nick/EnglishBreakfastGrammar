"""
export_enriched_corpus.py
Exports the fully enriched staging corpus from data/staging.db to:
1. Clean Excel CMS Workbooks:
   - data/cms/english_cms_final_enriched.xlsx
   - C:\\Users\\user\\Desktop\\english_cms_final_enriched.xlsx
2. Canonical Universal Lesson JSON:
   - data/cms/universal_lessons_final_enriched.json
   - C:\\Users\\user\\Desktop\\universal_lessons_final_enriched.json
"""

from pathlib import Path
import argparse
import json
import logging
import re
import sqlite3
import sys
from typing import Dict, Any, List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
DEFAULT_DB = REPO_ROOT / "data" / "staging.db"

LOCAL_CMS_XLSX = REPO_ROOT / "data" / "cms" / "english_cms_final_enriched.xlsx"
DESKTOP_CMS_XLSX = Path.home() / "Desktop" / "english_cms_final_enriched.xlsx"

LOCAL_UNIVERSAL_JSON = REPO_ROOT / "data" / "cms" / "universal_lessons_final_enriched.json"
DESKTOP_UNIVERSAL_JSON = Path.home() / "Desktop" / "universal_lessons_final_enriched.json"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s"
)
logger = logging.getLogger("export_enriched_corpus")


def setup_sheet(ws, headers: List[str]):
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


def finish_sheet(ws):
    body_alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                val = str(cell.value)
                if len(val) > max_length:
                    max_length = len(val)
        width = min(max(max_length + 2, 12), 50)
        ws.column_dimensions[column_letter].width = width

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = body_alignment


def normalize_placeholders(text: str) -> str:
    """Normalizes {{gap_1}} -> {{gap1}} on the export boundary to match canonical JS format."""
    if not text:
        return ""
    return re.sub(r"\{\{gap_(\d+)\}\}", r"{{gap\1}}", text)


def export_corpus_from_staging(
    db_path: Path = DEFAULT_DB,
    local_xlsx: Path = LOCAL_CMS_XLSX,
    desktop_xlsx: Path = DESKTOP_CMS_XLSX,
    local_json: Path = LOCAL_UNIVERSAL_JSON,
    desktop_json: Path = DESKTOP_UNIVERSAL_JSON,
) -> Dict[str, Any]:
    if not db_path.exists():
        raise FileNotFoundError(f"Staging database not found at: {db_path}")

    conn = sqlite3.connect(str(db_path))
    cursor = conn.cursor()

    # 1. Fetch all data relationally
    cursor.execute("""
        SELECT lesson_id, title, level, topic, status, description,
               source_provider, source_url, created_at, updated_at
        FROM staging_lessons
        ORDER BY lesson_id;
    """)
    lessons_rows = cursor.fetchall()

    cursor.execute("""
        SELECT exercise_id, lesson_id, page, exercise_order, title,
               instruction, example_source, example_target, status,
               source_file, source_url
        FROM staging_exercises
        ORDER BY lesson_id, page, exercise_order;
    """)
    exercises_rows = cursor.fetchall()

    cursor.execute("""
        SELECT question_id, exercise_id, lesson_id, question_order,
               response_model, content, explanation, difficulty, status
        FROM staging_questions
        ORDER BY exercise_id, question_order;
    """)
    questions_rows = cursor.fetchall()

    cursor.execute("""
        SELECT gap_id, question_id, gap_order, input_control,
               correct_answer, accepted_answers, case_sensitive,
               feedback_correct, feedback_incorrect
        FROM staging_gaps
        ORDER BY question_id, gap_order;
    """)
    gaps_rows = cursor.fetchall()

    cursor.execute("""
        SELECT option_id, question_id, gap_id, option_order,
               text, value, is_correct
        FROM staging_options
        ORDER BY question_id, gap_id, option_order;
    """)
    options_rows = cursor.fetchall()

    conn.close()

    logger.info(
        f"Fetched from DB: {len(lessons_rows)} lessons, {len(exercises_rows)} exercises, "
        f"{len(questions_rows)} questions, {len(gaps_rows)} gaps, {len(options_rows)} options."
    )

    # 2. Build CMS Workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_lessons = wb.create_sheet("Lessons")
    setup_sheet(ws_lessons, [
        "lesson_id", "title", "level", "topic", "status",
        "description", "source_provider", "source_url", "created_at", "updated_at"
    ])
    for r in lessons_rows:
        ws_lessons.append([
            r[0], r[1], r[2], r[3], "draft", r[5], r[6], r[7], r[8], r[9]
        ])

    ws_exercises = wb.create_sheet("Exercises")
    setup_sheet(ws_exercises, [
        "exercise_id", "lesson_id", "page", "order", "title",
        "instruction", "example_source", "example_target", "status", "source_file"
    ])
    for r in exercises_rows:
        ws_exercises.append([
            r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], "draft", r[9]
        ])

    ws_questions = wb.create_sheet("Questions")
    setup_sheet(ws_questions, [
        "question_id", "exercise_id", "lesson_id", "order",
        "response_model", "content", "explanation", "difficulty", "status"
    ])
    for r in questions_rows:
        ws_questions.append([
            r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], "draft"
        ])

    ws_gaps = wb.create_sheet("Gaps")
    setup_sheet(ws_gaps, [
        "gap_id", "question_id", "gap_order", "input_control",
        "correct_answer", "accepted_answers", "case_sensitive",
        "feedback_correct", "feedback_incorrect"
    ])
    for r in gaps_rows:
        # r[5] is accepted_answers JSON string e.g. '["answer"]'
        ws_gaps.append([
            r[0], r[1], r[2], r[3], r[4], r[5], bool(r[6]), r[7], r[8]
        ])

    ws_options = wb.create_sheet("Options")
    setup_sheet(ws_options, [
        "option_id", "question_id", "gap_id", "order",
        "text", "value", "is_correct"
    ])
    for r in options_rows:
        corr_bool = bool(r[6] == 1) if r[6] is not None else None
        ws_options.append([
            r[0], r[1], r[2] or "", r[3], r[4], r[5] or r[4], corr_bool
        ])

    ws_explanations = wb.create_sheet("Explanations")
    setup_sheet(ws_explanations, [
        "explanation_id", "lesson_id", "question_id", "section_order",
        "title", "content", "example", "status"
    ])

    finish_sheet(ws_lessons)
    finish_sheet(ws_exercises)
    finish_sheet(ws_questions)
    finish_sheet(ws_gaps)
    finish_sheet(ws_options)
    finish_sheet(ws_explanations)

    # Save CMS workbook
    local_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(local_xlsx))
    logger.info(f"Saved CMS Excel to local: {local_xlsx}")

    try:
        desktop_xlsx.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(desktop_xlsx))
        logger.info(f"Saved CMS Excel to desktop: {desktop_xlsx}")
    except Exception as e:
        logger.warning(f"Failed to save to desktop CMS Excel: {e}")

    # 3. Build Canonical Universal JSON
    # Organize hierarchically: options -> gaps -> questions -> exercises -> lessons
    options_by_gap = {}
    options_by_question = {}
    for r in options_rows:
        opt_id, q_id, gap_id, opt_order, text, val, is_corr = r
        corr_bool = bool(is_corr == 1) if is_corr is not None else None
        opt_obj = {
            "id": opt_id,
            "order": opt_order,
            "text": text,
            "value": val or text,
            "is_correct": corr_bool
        }
        if gap_id:
            options_by_gap.setdefault(gap_id, []).append(opt_obj)
        else:
            options_by_question.setdefault(q_id, []).append(opt_obj)

    gaps_by_question = {}
    for r in gaps_rows:
        gap_id, q_id, gap_order, input_ctrl, corr_ans, acc_json, case_sens, fb_corr, fb_inc = r
        try:
            acc_list = json.loads(acc_json) if acc_json else []
        except Exception:
            acc_list = [corr_ans] if corr_ans else []

        gap_opts = options_by_gap.get(gap_id, [])
        gap_obj = {
            "id": gap_id,
            "order": gap_order,
            "placeholder": f"{{{{gap{gap_order}}}}}",
            "input_control": input_ctrl,
            "correct_answer": corr_ans,
            "accepted_answers": acc_list,
            "options": gap_opts
        }
        gaps_by_question.setdefault(q_id, []).append(gap_obj)

    questions_by_exercise = {}
    for r in questions_rows:
        q_id, ex_id, l_id, q_order, resp_model, raw_content, expl, diff, status = r
        norm_text = normalize_placeholders(raw_content)
        q_gaps = gaps_by_question.get(q_id, [])
        q_opts = options_by_question.get(q_id, [])

        if resp_model == "gap" or len(q_gaps) > 0:
            first_ctrl = q_gaps[0].get("input_control", "").lower() if q_gaps else ""
            q_type = "gap_text" if first_ctrl == "text" else "gap_select"
        else:
            q_type = resp_model

        q_obj = {
            "id": q_id,
            "order": q_order,
            "type": q_type,
            "text": norm_text,
            "explanation": expl or "",
            "difficulty": diff or "",
            "status": "draft"
        }
        if q_gaps:
            q_obj["gaps"] = q_gaps
        if q_opts:
            q_obj["options"] = q_opts

        questions_by_exercise.setdefault(ex_id, []).append(q_obj)

    exercises_by_lesson = {}
    for r in exercises_rows:
        ex_id, l_id, page, ex_order, title, instr, ex_src, ex_tgt, status, src_file, src_url = r
        ex_questions = questions_by_exercise.get(ex_id, [])
        ex_obj = {
            "id": ex_id,
            "lessonId": l_id,
            "order": ex_order,
            "page": page,
            "title": title or "",
            "instruction": instr or "",
            "status": "draft",
            "source_file": src_file or "",
            "questions": ex_questions
        }
        exercises_by_lesson.setdefault(l_id, []).append(ex_obj)

    lessons_list = []
    for r in lessons_rows:
        l_id, title, level, topic, status, desc, provider, src_url, created, updated = r
        lesson_exercises = exercises_by_lesson.get(l_id, [])
        lesson_obj = {
            "id": l_id,
            "title": title,
            "level": level,
            "topic": topic,
            "status": "draft",
            "description": desc or "",
            "order": 1,
            "version": 1,
            "exercises": lesson_exercises,
            "source": {
                "provider": provider,
                "source_url": src_url or ""
            }
        }
        lessons_list.append(lesson_obj)

    universal_dataset = {
        "version": "1.0",
        "provider": "Test-English",
        "total_lessons": len(lessons_list),
        "total_exercises": len(exercises_rows),
        "total_questions": len(questions_rows),
        "total_gaps": len(gaps_rows),
        "total_options": len(options_rows),
        "lessons": lessons_list
    }

    # Save Universal JSON
    local_json.parent.mkdir(parents=True, exist_ok=True)
    with open(local_json, "w", encoding="utf-8") as f:
        json.dump(universal_dataset, f, ensure_ascii=False, indent=2)
    logger.info(f"Saved Universal JSON to local: {local_json}")

    try:
        desktop_json.parent.mkdir(parents=True, exist_ok=True)
        with open(desktop_json, "w", encoding="utf-8") as f:
            json.dump(universal_dataset, f, ensure_ascii=False, indent=2)
        logger.info(f"Saved Universal JSON to desktop: {desktop_json}")
    except Exception as e:
        logger.warning(f"Failed to save to desktop Universal JSON: {e}")

    report = {
        "lessons_exported": len(lessons_rows),
        "exercises_exported": len(exercises_rows),
        "questions_exported": len(questions_rows),
        "gaps_exported": len(gaps_rows),
        "options_exported": len(options_rows),
        "local_xlsx": str(local_xlsx),
        "desktop_xlsx": str(desktop_xlsx),
        "local_json": str(local_json),
        "desktop_json": str(desktop_json),
    }
    return report


def main():
    parser = argparse.ArgumentParser(description="Export fully enriched staging corpus.")
    parser.add_argument("--db", type=Path, default=DEFAULT_DB, help="Path to staging.db")
    args = parser.parse_args()

    report = export_corpus_from_staging(db_path=args.db)

    print("\n" + "=" * 60)
    print("ENRICHED CORPUS EXPORT REPORT")
    print("=" * 60)
    for k, v in report.items():
        print(f"  {k}: {v}")
    print("=" * 60)


if __name__ == "__main__":
    main()
