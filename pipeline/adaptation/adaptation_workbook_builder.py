"""Adaptation Workbook Builder (TASK-013).

Universal English Test Platform
Generates the Google Sheets / Gemini adaptation workbook for all 5,571 PENDING
questions in data/adaptation.db.

STRICT CONTRACT:
1. Does NOT modify staging.db (immutable source).
2. Does NOT modify data/adaptation.db.
3. Excludes all 225 already completed/validated questions (200 pilot + 25 dry-run).
4. Strictly orders all 5,571 questions level-by-level:
   A1 -> A2 -> B1 -> B1-B2 -> B2 -> C1 -> Shorts
5. Enforces the Core Adaptation Principle:
   Preserve the source correct answer whenever reasonably possible.
6. Populates model-tailored Google Sheets =IFERROR(GEMINI(...), "") formulas.
7. Writes to:
   - data/adaptation/english_adaptation_gemini_5571.xlsx
   - C:\\Users\\user\\Desktop\\english_adaptation_gemini_5571.xlsx
"""

from __future__ import annotations

import argparse
import datetime
import logging
from pathlib import Path
import sqlite3
import sys
from typing import Any, Dict, List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

DEFAULT_ADAPTATION_DB = REPO_ROOT / "data" / "adaptation.db"
DEFAULT_STAGING_DB = REPO_ROOT / "data" / "staging.db"
LOCAL_OUTPUT = REPO_ROOT / "data" / "adaptation" / "english_adaptation_gemini_5571.xlsx"
DESKTOP_OUTPUT = Path.home() / "Desktop" / "english_adaptation_gemini_5571.xlsx"

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("adaptation_workbook_builder")

LEVEL_ORDER = {
    "A1": 1,
    "A2": 2,
    "B1": 3,
    "B1-B2": 4,
    "B2": 5,
    "C1": 6,
    "SHORTS": 7,
}


def fetch_pending_adaptation_questions(
    adaptation_db_path: Path, staging_db_path: Path
) -> List[Dict[str, Any]]:
    """Fetches all 5,571 PENDING questions joined with immutable staging metadata, ordered level-by-level."""
    adapt_conn = sqlite3.connect(f"file:{adaptation_db_path.resolve()}?mode=ro", uri=True)
    adapt_conn.row_factory = sqlite3.Row
    stage_conn = sqlite3.connect(f"file:{staging_db_path.resolve()}?mode=ro", uri=True)
    stage_conn.row_factory = sqlite3.Row

    try:
        # Fetch pending questions from adaptation.db
        pending_rows = adapt_conn.execute(
            """
            SELECT adapted_question_id, source_question_id, adapted_exercise_id,
                   adapted_lesson_id, question_order, response_model, source_text
            FROM adapted_questions
            WHERE adaptation_status = 'PENDING'
            """
        ).fetchall()

        logger.info(f"Found {len(pending_rows)} PENDING questions in adaptation.db")
        if len(pending_rows) != 5571:
            logger.warning(f"Expected 5,571 PENDING questions, found {len(pending_rows)}")

        items: List[Dict[str, Any]] = []

        for row in pending_rows:
            sqid = str(row["source_question_id"])
            aqid = str(row["adapted_question_id"])
            rm = str(row["response_model"])

            # Staging metadata
            s_q = stage_conn.execute("SELECT * FROM staging_questions WHERE question_id = ?", (sqid,)).fetchone()
            s_ex = stage_conn.execute("SELECT * FROM staging_exercises WHERE exercise_id = ?", (s_q["exercise_id"],)).fetchone()
            s_les = stage_conn.execute("SELECT * FROM staging_lessons WHERE lesson_id = ?", (s_q["lesson_id"],)).fetchone()

            level = s_les["level"].upper()
            topic = s_les["topic"]
            lesson_title = s_les["title"]
            ex_title = s_ex["title"]
            instruction = (s_ex["instruction"] or "").strip()

            # Options context
            s_opts = stage_conn.execute(
                "SELECT option_order, text, is_correct FROM staging_options WHERE question_id = ? ORDER BY option_order", (sqid,)
            ).fetchall()
            # Gaps context
            s_gaps = stage_conn.execute(
                "SELECT gap_order, correct_answer, accepted_answers FROM staging_gaps WHERE question_id = ? ORDER BY gap_order", (sqid,)
            ).fetchall()

            opts_context_parts = []
            target_answers_parts = []
            if rm in ("single_choice", "multiple_choice"):
                for o in s_opts:
                    tag = " [CORRECT]" if o["is_correct"] else ""
                    opts_context_parts.append(f"[{o['option_order']}] {o['text']}{tag}")
                    if o["is_correct"]:
                        target_answers_parts.append(str(o["text"]))
            elif rm == "gap":
                for g in s_gaps:
                    opts_context_parts.append(f"Gap {g['gap_order']}: target='{g['correct_answer']}'")
                    target_answers_parts.append(str(g["correct_answer"]))
                if s_opts:
                    opt_list = " | ".join([str(o["text"]) for o in s_opts])
                    opts_context_parts.append(f"Select choices: {opt_list}")

            opts_context_str = " | ".join(opts_context_parts)
            target_answer_str = " | ".join(target_answers_parts)

            items.append({
                "source_question_id": sqid,
                "adapted_question_id": aqid,
                "lesson_id": s_q["lesson_id"],
                "exercise_id": s_q["exercise_id"],
                "level": level,
                "topic": topic,
                "lesson_title": lesson_title,
                "exercise_title": ex_title,
                "question_order": s_q["question_order"],
                "response_model": rm,
                "source_text": s_q["content"],
                "instruction": instruction,
                "options_context": opts_context_str,
                "target_answer": target_answer_str,
                "num_options": len(s_opts),
                "num_gaps": len(s_gaps),
            })

    finally:
        adapt_conn.close()
        stage_conn.close()

    # Sort strictly level-by-level: A1 -> A2 -> B1 -> B1-B2 -> B2 -> C1 -> Shorts
    items.sort(
        key=lambda x: (
            LEVEL_ORDER.get(x["level"], 99),
            x["lesson_id"],
            x["exercise_id"],
            x["question_order"],
        )
    )

    return items


def build_gemini_formula(row_idx: int, response_model: str) -> str:
    """Builds model-tailored Google Sheets AI/Gemini formula for a specific row index."""
    # Col D: level, Col E: topic, Col F: instruction, Col G: source_text, Col H: options_context, Col I: target_answer
    if response_model == "gap":
        prompt = (
            f'"Task: Adapt English test question for Level " & D{row_idx} & " (" & E{row_idx} & "). '
            f'Instruction: " & F{row_idx} & ". Source sentence: \'" & G{row_idx} & "\'. '
            f'Gaps: " & H{row_idx} & ". TARGET ANSWER(S) TO PRESERVE: \'" & I{row_idx} & "\'. '
            'RULES: '
            '1. PRESERVE THE TARGET ANSWER: The adapted sentence MUST make the exact same target answer(s) correct. '
            '2. REWRITE CONTEXT: Completely change the situation, names, setting, and surrounding words. '
            '3. ZERO PLAGIARISM: Do not reuse 3+ consecutive words from source text. '
            '4. Keep the exact same number of {{gap_1}} placeholders. '
            'Return ONLY valid JSON: {""adapted_text"":""sentence with {{gap_1}}..."",""gaps"":[{""gap_order"":1,""correct_answer"":""...""}],""target_tokens"":[""...""]}. '
            'Output strictly JSON without markdown or explanations."'
        )
    elif response_model == "single_choice":
        prompt = (
            f'"Task: Adapt single-choice English question for Level " & D{row_idx} & " (" & E{row_idx} & "). '
            f'Instruction: " & F{row_idx} & ". Source sentence: \'" & G{row_idx} & "\'. '
            f'Options: " & H{row_idx} & ". TARGET CORRECT ANSWER: \'" & I{row_idx} & "\'. '
            'RULES: '
            '1. PRESERVE CORRECT ANSWER: The adapted question MUST preserve the exact same correct option and grammatical error types for distractors. '
            '2. REWRITE SCENARIO: Completely change the context, names, objects, and situation. '
            '3. ZERO PLAGIARISM: Do not reuse 3+ consecutive non-target words. '
            '4. Return ONLY valid JSON: {""adapted_text"":""new sentence with _____ blank"",""options"":[{""text"":""..."",""is_correct"":0},{""text"":""..."",""is_correct"":1},{""text"":""..."",""is_correct"":0}],""target_tokens"":[""...""]}. '
            'Output strictly JSON without markdown or explanations."'
        )
    elif response_model == "multiple_choice":
        prompt = (
            f'"Task: Adapt multiple-choice English question for Level " & D{row_idx} & " (" & E{row_idx} & "). '
            f'Instruction: " & F{row_idx} & ". Source sentence: \'" & G{row_idx} & "\'. '
            f'Options: " & H{row_idx} & ". TARGET CORRECT ANSWERS: \'" & I{row_idx} & "\'. '
            'RULES: '
            '1. PRESERVE CORRECT ANSWERS: The adapted question MUST preserve the same TWO correct answers. '
            '2. REWRITE SCENARIO: Completely change narrative context, names, and environment. '
            '3. ZERO PLAGIARISM: Do not reuse 3+ consecutive words. '
            '4. Return ONLY valid JSON: {""adapted_text"":""sentence with blank... Choose TWO correct answers"",""options"":[{""text"":""..."",""is_correct"":1},{""text"":""..."",""is_correct"":1},{""text"":""..."",""is_correct"":0},{""text"":""..."",""is_correct"":0}],""target_tokens"":[""...""]}. '
            'Output strictly JSON without markdown or explanations."'
        )
    else:
        prompt = (
            f'"Task: Adapt English test question. Source: \'" & G{row_idx} & "\'. '
            f'Target Answer: \'" & I{row_idx} & "\'. Return ONLY valid JSON adaptation."'
        )

    return f'=IFERROR(GEMINI({prompt}), "")'


def create_adaptation_workbook(
    items: List[Dict[str, Any]],
    output_path: Path,
    desktop_path: Optional[Path] = None,
) -> Dict[str, Any]:
    """Builds the 2-sheet adaptation workbook (Summary + Gemini_Adaptation)."""
    wb = openpyxl.Workbook()
    # Default active sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"

    # Styling definitions
    FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    FONT_TITLE = Font(name="Calibri", size=15, bold=True, color="1F4E78")
    FONT_SECTION = Font(name="Calibri", size=12, bold=True, color="1F4E78")
    FONT_BOLD = Font(name="Calibri", size=11, bold=True)
    FONT_REGULAR = Font(name="Calibri", size=11)
    FILL_HEADER = PatternFill(fill_type="solid", fgColor="1F4E78")
    FILL_SECTION = PatternFill(fill_type="solid", fgColor="D9E1F2")
    FILL_ACCENT = PatternFill(fill_type="solid", fgColor="E2EFDA")
    BORDER_THIN = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )

    # =========================================================================
    # SHEET 1: SUMMARY
    # =========================================================================
    ws_summary.views.sheetView[0].showGridLines = True
    ws_summary.cell(row=1, column=2, value="UNIVERSAL ENGLISH TEST PLATFORM — FULL CORPUS ADAPTATION WORKBOOK").font = FONT_TITLE
    ws_summary.cell(row=2, column=2, value="TASK-013: Production Batch Generation for 5,571 PENDING Questions").font = FONT_SECTION
    ws_summary.row_dimensions[1].height = 28
    ws_summary.row_dimensions[2].height = 22

    # Level breakdown
    level_counts: Dict[str, int] = {}
    model_counts: Dict[str, int] = {}
    for it in items:
        level_counts[it["level"]] = level_counts.get(it["level"], 0) + 1
        model_counts[it["response_model"]] = model_counts.get(it["response_model"], 0) + 1

    ws_summary.cell(row=4, column=2, value="1. EXECUTION LEVEL ORDER (STRICT SEQUENTIAL DISPATCH)").font = FONT_SECTION
    sum_headers = ["Order", "CEFR Level", "Pending Questions", "Proportion", "Target Model Distribution"]
    for c_idx, h in enumerate(sum_headers, start=2):
        cell = ws_summary.cell(row=5, column=c_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[5].height = 24

    cur_row = 6
    for lvl in ["A1", "A2", "B1", "B1-B2", "B2", "C1", "SHORTS"]:
        cnt = level_counts.get(lvl, 0)
        pct = cnt / len(items) if items else 0.0
        ws_summary.cell(row=cur_row, column=2, value=LEVEL_ORDER.get(lvl, 99)).alignment = Alignment(horizontal="center")
        ws_summary.cell(row=cur_row, column=3, value=lvl).font = FONT_BOLD
        ws_summary.cell(row=cur_row, column=4, value=cnt).alignment = Alignment(horizontal="center")
        ws_summary.cell(row=cur_row, column=5, value=f"{pct:.1%}").alignment = Alignment(horizontal="center")
        ws_summary.cell(row=cur_row, column=6, value="gap + single_choice + multiple_choice")
        for col in range(2, 7):
            ws_summary.cell(row=cur_row, column=col).border = BORDER_THIN
        cur_row += 1

    # Total row
    ws_summary.cell(row=cur_row, column=3, value="TOTAL PENDING").font = FONT_BOLD
    ws_summary.cell(row=cur_row, column=4, value=len(items)).font = FONT_BOLD
    ws_summary.cell(row=cur_row, column=4).alignment = Alignment(horizontal="center")
    ws_summary.cell(row=cur_row, column=5, value="100.0%").font = FONT_BOLD
    ws_summary.cell(row=cur_row, column=5).alignment = Alignment(horizontal="center")
    for col in range(2, 7):
        ws_summary.cell(row=cur_row, column=col).fill = FILL_SECTION
        ws_summary.cell(row=cur_row, column=col).border = BORDER_THIN
    cur_row += 2

    # Instructions box
    ws_summary.cell(row=cur_row, column=2, value="2. GOOGLE SHEETS EVALUATION INSTRUCTIONS").font = FONT_SECTION
    cur_row += 1
    instructions = [
        ("Step 1: Upload to Google Drive", "Upload this workbook to Google Drive and open with Google Sheets."),
        ("Step 2: Enable Gemini", "Ensure the Google Workspace Gemini extension or =GEMINI() function is enabled."),
        ("Step 3: Evaluate Batch-by-Batch", "Select Column J (gemini_formula) by level (A1 first, then A2, etc.) to trigger generation."),
        ("Step 4: Copy Evaluated Values", "Once generated, copy Column J and 'Paste special -> Values only' into Column K (gemini_raw_values)."),
        ("Step 5: Download & Import", "Download as XLSX and run: npm run adaptation:import-results"),
    ]
    for title, desc in instructions:
        ws_summary.cell(row=cur_row, column=2, value=title).font = FONT_BOLD
        c_desc = ws_summary.cell(row=cur_row, column=3, value=desc)
        c_desc.font = FONT_REGULAR
        ws_summary.merge_cells(start_row=cur_row, start_column=3, end_row=cur_row, end_column=7)
        cur_row += 1

    ws_summary.column_dimensions["B"].width = 30
    ws_summary.column_dimensions["C"].width = 16
    ws_summary.column_dimensions["D"].width = 20
    ws_summary.column_dimensions["E"].width = 16
    ws_summary.column_dimensions["F"].width = 36
    ws_summary.column_dimensions["G"].width = 24

    # =========================================================================
    # SHEET 2: GEMINI_ADAPTATION
    # =========================================================================
    ws_adapt = wb.create_sheet("Gemini_Adaptation")
    ws_adapt.views.sheetView[0].showGridLines = True
    ws_adapt.freeze_panes = "A3"

    # Title row
    title_c = ws_adapt.cell(row=1, column=1, value="ENGLISH ADAPTATION CORPUS — 5,571 PENDING QUESTIONS (TASK-013)")
    title_c.font = Font(name="Calibri", size=13, bold=True, color="1F4E78")
    ws_adapt.row_dimensions[1].height = 24

    headers = [
        "source_question_id",     # Col A (1)
        "adapted_question_id",    # Col B (2)
        "exercise_id",            # Col C (3)
        "level",                  # Col D (4)
        "topic",                  # Col E (5)
        "instruction",            # Col F (6)
        "source_sentence",        # Col G (7)
        "options_context",        # Col H (8)
        "target_answer_preserve", # Col I (9)
        "gemini_formula",         # Col J (10)
        "gemini_raw_values",      # Col K (11)
        "response_model",         # Col L (12)
        "lesson_id",              # Col M (13)
        "question_order",         # Col N (14)
    ]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws_adapt.cell(row=2, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_adapt.row_dimensions[2].height = 28

    # Populate rows
    logger.info(f"Populating {len(items)} questions into Gemini_Adaptation sheet...")
    for idx, it in enumerate(items, start=3):
        sqid = it["source_question_id"]
        aqid = it["adapted_question_id"]
        eid = it["exercise_id"]
        lvl = it["level"]
        topic = it["topic"]
        inst = it["instruction"]
        stext = it["source_text"]
        opts_ctx = it["options_context"]
        tgt_ans = it["target_answer"]
        rm = it["response_model"]
        lid = it["lesson_id"]
        qord = it["question_order"]

        formula = build_gemini_formula(idx, rm)

        ws_adapt.cell(row=idx, column=1, value=int(sqid))
        ws_adapt.cell(row=idx, column=2, value=aqid)
        ws_adapt.cell(row=idx, column=3, value=eid)
        ws_adapt.cell(row=idx, column=4, value=lvl)
        ws_adapt.cell(row=idx, column=5, value=topic)
        ws_adapt.cell(row=idx, column=6, value=inst)
        ws_adapt.cell(row=idx, column=7, value=stext)
        ws_adapt.cell(row=idx, column=8, value=opts_ctx)
        ws_adapt.cell(row=idx, column=9, value=tgt_ans)
        ws_adapt.cell(row=idx, column=10, value=formula)
        ws_adapt.cell(row=idx, column=11, value="")  # To be filled by user
        ws_adapt.cell(row=idx, column=12, value=rm)
        ws_adapt.cell(row=idx, column=13, value=lid)
        ws_adapt.cell(row=idx, column=14, value=qord)

        # Basic cell styling
        for col_i in range(1, 15):
            c = ws_adapt.cell(row=idx, column=col_i)
            c.border = BORDER_THIN
            c.font = FONT_REGULAR
            if col_i in (1, 4, 12, 14):
                c.alignment = Alignment(horizontal="center", vertical="top")
            elif col_i in (2, 3, 5, 13):
                c.alignment = Alignment(horizontal="left", vertical="top")
            else:
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        ws_adapt.row_dimensions[idx].height = 32

    # Column widths
    col_widths = {
        1: 18,  # source_question_id
        2: 20,  # adapted_question_id
        3: 14,  # exercise_id
        4: 10,  # level
        5: 28,  # topic
        6: 35,  # instruction
        7: 45,  # source_sentence
        8: 40,  # options_context
        9: 25,  # target_answer_preserve
        10: 45, # gemini_formula
        11: 45, # gemini_raw_values
        12: 16, # response_model
        13: 28, # lesson_id
        14: 12, # question_order
    }
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws_adapt.column_dimensions[col_letter].width = width

    # Save to local path
    output_path.parent.mkdir(parents=True, exist_ok=True)
    logger.info(f"Saving workbook to {output_path}...")
    wb.save(output_path)
    logger.info(f"Saved local copy: {output_path} ({output_path.stat().st_size:,} bytes)")

    # Save to desktop path if provided
    if desktop_path:
        try:
            logger.info(f"Saving desktop copy to {desktop_path}...")
            wb.save(desktop_path)
            logger.info(f"Saved desktop copy: {desktop_path} ({desktop_path.stat().st_size:,} bytes)")
        except Exception as e:
            logger.warning(f"Failed to save desktop copy: {e}")

    return {
        "total_items": len(items),
        "level_breakdown": level_counts,
        "model_breakdown": model_counts,
        "local_path": str(output_path),
        "desktop_path": str(desktop_path) if desktop_path else None,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Google Sheets / Gemini Adaptation Workbook (TASK-013)")
    parser.add_argument("--adapt-db", default=str(DEFAULT_ADAPTATION_DB), help="Path to adaptation.db")
    parser.add_argument("--stage-db", default=str(DEFAULT_STAGING_DB), help="Path to staging.db")
    parser.add_argument("--out", default=str(LOCAL_OUTPUT), help="Local output path")
    parser.add_argument("--desktop", default=str(DESKTOP_OUTPUT), help="Desktop output path")
    args = parser.parse_args()

    print("=" * 65)
    print(" ADAPTATION WORKBOOK BUILDER (TASK-013)")
    print("=" * 65)

    items = fetch_pending_adaptation_questions(Path(args.adapt_db), Path(args.stage_db))
    res = create_adaptation_workbook(items, Path(args.out), Path(args.desktop))

    print(f"\nTotal Pending Questions : {res['total_items']}")
    print(f"Level Breakdown         : {res['level_breakdown']}")
    print(f"Model Breakdown         : {res['model_breakdown']}")
    print(f"Local Output            : {res['local_path']}")
    print(f"Desktop Output          : {res['desktop_path']}")
    print("\n[SUCCESS] Adaptation workbook created successfully!")


if __name__ == "__main__":
    main()
