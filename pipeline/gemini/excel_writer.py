from __future__ import annotations

from openpyxl.workbook import Workbook

from models import ValidationResult


# ============================================================
# EXCEL WRITER
# ============================================================

def write_results(
    result:     ValidationResult,
    wb:         Workbook,
    logger,
) -> tuple[int, int]:
    """
    Writes validated answers into the workbook (in memory).
    The caller is responsible for saving the file.

    Writes to:
      - Gaps sheet:   correct_answer, accepted_answers
      - Options sheet: is_correct

    Never modifies: question_id, gap_id, option_id, text,
                    content, instruction, or any other field.

    Returns (gaps_written, options_written) counts.
    """

    ws_gaps    = wb["Gaps"]
    ws_options = wb["Options"]

    # --------------------------------------------------------
    # Build row-index maps: id → row_number (1-based)
    # --------------------------------------------------------

    gaps_header, gap_rows       = _index_sheet(ws_gaps,    "gap_id")
    options_header, option_rows = _index_sheet(ws_options, "option_id")

    # --------------------------------------------------------
    # Column positions
    # --------------------------------------------------------

    gap_col_correct   = _col(gaps_header, "correct_answer")
    gap_col_accepted  = _col(gaps_header, "accepted_answers")
    opt_col_correct   = _col(options_header, "is_correct")

    gaps_written    = 0
    options_written = 0

    # --------------------------------------------------------
    # Write per validated question
    # --------------------------------------------------------

    for vq in result.questions:

        # ---- GAPS ----
        for vg in vq.gaps:

            row = gap_rows.get(vg.gap_id)

            if row is None:
                logger.warning(
                    f"[WRITE] gap_id '{vg.gap_id}' "
                    f"not found in Gaps sheet — skipping"
                )
                continue

            # accepted_answers: store extras only (not correct_answer)
            # JSON Builder will reconstruct full list.
            # E.g. correct="is", accepted=["is","'s"] → extras="'s"
            extras = [
                a for a in vg.accepted_answers
                if a != vg.correct_answer
            ]
            accepted_str = " | ".join(extras) if extras else ""

            ws_gaps.cell(row=row, column=gap_col_correct).value  = vg.correct_answer
            ws_gaps.cell(row=row, column=gap_col_accepted).value = accepted_str

            logger.info(
                f"[GAP WRITE] {vg.gap_id}: "
                f"correct_answer='{vg.correct_answer}' "
                f"accepted_answers='{accepted_str}'"
            )
            gaps_written += 1

        # ---- OPTIONS ----
        for vo in vq.options:

            row = option_rows.get(vo.option_id)

            if row is None:
                logger.warning(
                    f"[WRITE] option_id '{vo.option_id}' "
                    f"not found in Options sheet — skipping"
                )
                continue

            ws_options.cell(
                row=row,
                column=opt_col_correct,
            ).value = vo.is_correct

            logger.debug(
                f"[OPT WRITE] {vo.option_id}: is_correct={vo.is_correct}"
            )
            options_written += 1

    return gaps_written, options_written


# ============================================================
# HELPERS
# ============================================================

def _index_sheet(
    ws,
    id_column: str,
) -> tuple[list[str], dict[str, int]]:
    """
    Reads headers from row 1 and builds a dict: id_value → row_number.
    Row numbers are 1-based (openpyxl convention).
    """

    rows    = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h else "" for h in rows[0]]

    id_col_idx = headers.index(id_column)

    id_to_row: dict[str, int] = {}

    for row_idx, row in enumerate(rows[1:], start=2):
        cell_val = row[id_col_idx]
        if cell_val is not None:
            id_to_row[str(cell_val).strip()] = row_idx

    return headers, id_to_row


def _col(headers: list[str], name: str) -> int:
    """Returns 1-based column number for the given header name."""
    return headers.index(name) + 1
