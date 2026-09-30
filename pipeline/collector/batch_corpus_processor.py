"""
batch_corpus_processor.py
Batch offline processor for the entire cached Test-English HTML corpus.

Parses all 638 exercise pages across 225 topics and 7 levels into:
1. data/cms/universal_lessons_preliminary.json (and Desktop copy)
2. data/cms/english_cms_preliminary.xlsx (and Desktop copy)

STRICT SAFETY:
Never modifies or overwrites english_cms.xlsx or english_cms_gemini_all_182.xlsx.
"""

from pathlib import Path
import json
import logging
import re
import sys
from typing import Dict, List, Any

# openpyxl for Excel CMS generation
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Ensure pipeline root is accessible
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from parser.test_english_parser import parse_lesson_from_html, clean_text

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s"
)
logger = logging.getLogger("batch_corpus_processor")

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
MANIFEST_PATH = REPO_ROOT / "data" / "cache" / "acquisition_manifest.json"
CATALOG_PATH = REPO_ROOT / "data" / "catalog" / "source_catalog.json"
CACHE_DIR = REPO_ROOT / "data" / "cache" / "html"

OUTPUT_JSON_LOCAL = REPO_ROOT / "data" / "cms" / "universal_lessons_preliminary.json"
OUTPUT_XLSX_LOCAL = REPO_ROOT / "data" / "cms" / "english_cms_preliminary.xlsx"

DESKTOP_DIR = Path.home() / "Desktop"
OUTPUT_JSON_DESKTOP = DESKTOP_DIR / "universal_lessons_preliminary.json"
OUTPUT_XLSX_DESKTOP = DESKTOP_DIR / "english_cms_preliminary.xlsx"


def setup_sheet(ws, headers: List[str]):
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


def finish_sheet(ws):
    body_alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                val = str(cell.value)
                if len(val) > max_length:
                    max_length = len(val)
        width = min(max(max_length + 2, 12), 50)
        ws.column_dimensions[column_letter].width = width

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = body_alignment


def generate_cms_workbook(data: dict, output_path: Path):
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    # 1. Lessons
    ws_lessons = wb.create_sheet("Lessons")
    setup_sheet(ws_lessons, [
        "lesson_id", "title", "level", "topic", "status",
        "description", "source_provider", "source_url", "created_at", "updated_at"
    ])

    # 2. Exercises
    ws_exercises = wb.create_sheet("Exercises")
    setup_sheet(ws_exercises, [
        "exercise_id", "lesson_id", "page", "order", "title",
        "instruction", "example_source", "example_target", "status", "source_file"
    ])

    # 3. Questions
    ws_questions = wb.create_sheet("Questions")
    setup_sheet(ws_questions, [
        "question_id", "exercise_id", "lesson_id", "order",
        "response_model", "content", "explanation", "difficulty", "status"
    ])

    # 4. Gaps
    ws_gaps = wb.create_sheet("Gaps")
    setup_sheet(ws_gaps, [
        "gap_id", "question_id", "gap_order", "input_control",
        "correct_answer", "accepted_answers", "case_sensitive",
        "feedback_correct", "feedback_incorrect"
    ])

    # 5. Options
    ws_options = wb.create_sheet("Options")
    setup_sheet(ws_options, [
        "option_id", "question_id", "gap_id", "order",
        "text", "value", "is_correct"
    ])

    # 6. Explanations
    ws_explanations = wb.create_sheet("Explanations")
    setup_sheet(ws_explanations, [
        "explanation_id", "lesson_id", "question_id", "section_order",
        "title", "content", "example", "status"
    ])

    lessons = data.get("lessons", [])
    for lesson in lessons:
        lesson_id = lesson.get("lesson_id", "")
        ws_lessons.append([
            lesson_id,
            lesson.get("title", ""),
            lesson.get("level", ""),
            lesson.get("topic", ""),
            lesson.get("status", "draft"),
            lesson.get("description", ""),
            lesson.get("source", {}).get("provider", "Test-English"),
            lesson.get("source", {}).get("source_url", ""),
            "",
            ""
        ])

        pages = lesson.get("pages", [])
        for page in pages:
            page_number = page.get("page", 1)
            source_file = page.get("source_file", "")
            exercises = page.get("exercises", [])

            for exercise in exercises:
                exercise_id = exercise.get("exercise_id", "")
                ex_example = exercise.get("example") or {}
                ws_exercises.append([
                    exercise_id,
                    lesson_id,
                    page_number,
                    exercise.get("order", 1),
                    exercise.get("title", ""),
                    exercise.get("instruction", ""),
                    ex_example.get("source", ""),
                    ex_example.get("target", ""),
                    "draft",
                    source_file
                ])

                questions = exercise.get("questions", [])
                for q_idx, q in enumerate(questions, start=1):
                    question_id = q.get("question_id", "")
                    response_model = q.get("response_model", "")
                    content = q.get("content", "")

                    ws_questions.append([
                        question_id,
                        exercise_id,
                        lesson_id,
                        q.get("order", q_idx),
                        response_model,
                        content,
                        "",
                        "",
                        "draft"
                    ])

                    # Question-level options (single_choice, multiple_choice)
                    for opt_idx, opt in enumerate(q.get("options", []), start=1):
                        ws_options.append([
                            opt.get("option_id", f"{question_id}_{opt_idx}"),
                            question_id,
                            "",
                            opt.get("order", opt_idx),
                            opt.get("text", ""),
                            opt.get("value", ""),
                            opt.get("is_correct")  # None
                        ])

                    # Gaps
                    for gap_idx, gap in enumerate(q.get("gaps", []), start=1):
                        gap_id = gap.get("gap_id", f"gap_{gap_idx}")
                        ws_gaps.append([
                            gap_id,
                            question_id,
                            gap.get("gap_order", gap_idx),
                            gap.get("input_control", "text"),
                            gap.get("correct_answer"),  # None
                            json.dumps(gap.get("accepted_answers", [])),
                            False,
                            "",
                            ""
                        ])

                        # Gap options (select)
                        for g_opt_idx, g_opt in enumerate(gap.get("options", []), start=1):
                            ws_options.append([
                                g_opt.get("option_id", f"{gap_id}_{g_opt_idx}"),
                                question_id,
                                gap_id,
                                g_opt.get("order", g_opt_idx),
                                g_opt.get("text", ""),
                                g_opt.get("value", ""),
                                g_opt.get("is_correct")  # None
                            ])

    finish_sheet(ws_lessons)
    finish_sheet(ws_exercises)
    finish_sheet(ws_questions)
    finish_sheet(ws_gaps)
    finish_sheet(ws_options)
    finish_sheet(ws_explanations)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"Excel CMS workbook saved: {output_path}")


def process_entire_corpus() -> dict:
    logger.info("Loading acquisition manifest and source catalog...")
    with open(MANIFEST_PATH, "r", encoding="utf-8") as f:
        manifest = json.load(f)

    with open(CATALOG_PATH, "r", encoding="utf-8") as f:
        catalog = json.load(f)

    manifest_pages = manifest.get("pages", {})

    # Build URL to local cache path map
    url_to_cache = {}
    for url, info in manifest_pages.items():
        if info.get("topic_id") != "category_page":
            clean_url = url.rstrip("/")
            url_to_cache[clean_url] = info

    total_topics_catalog = 0
    lessons_list = []
    
    parsed_pages_count = 0
    total_exercises_count = 0
    total_questions_count = 0
    total_gaps_count = 0
    total_options_count = 0
    unresolved_answers_count = 0
    parser_errors = []
    parser_warnings = []
    topics_represented = set()

    for level_entry in catalog.get("levels", []):
        level_id = level_entry.get("level_id", "").upper()
        topics = level_entry.get("topics", [])

        for topic in topics:
            total_topics_catalog += 1
            topic_id = topic.get("topic_id", "")
            topic_title = topic.get("title", "")
            topic_url = topic.get("url", "").rstrip("/")
            seed_url = topic.get("url", "")

            lesson_id = f"{level_id.lower()}_{topic_id.replace('-', '_')}"
            topics_represented.add((level_id, topic_id))

            # Collect all page URLs for this topic
            exercise_urls = [topic_url]
            # Check pagination from manifest or catalog exercises
            for u in sorted(url_to_cache.keys()):
                if u.startswith(topic_url) and u != topic_url:
                    suffix = u[len(topic_url):].strip("/")
                    if suffix.isdigit():
                        exercise_urls.append(u)

            # Sort pages numerically
            def page_sort_key(u: str) -> int:
                if u == topic_url:
                    return 1
                suffix = u[len(topic_url):].strip("/")
                return int(suffix) if suffix.isdigit() else 999

            exercise_urls = sorted(set(exercise_urls), key=page_sort_key)

            lesson_pages = []
            exercise_global_order = 1

            for p_order, page_url in enumerate(exercise_urls, start=1):
                page_info = url_to_cache.get(page_url)
                if not page_info:
                    msg = f"Missing cache info for {page_url} in topic {topic_id} ({level_id})"
                    parser_warnings.append(msg)
                    logger.warning(msg)
                    continue

                local_path_str = page_info.get("local_cache_path")
                fpath = REPO_ROOT / local_path_str if local_path_str else None
                if not fpath or not fpath.exists():
                    # Try resolving directly in CACHE_DIR
                    clean_name = re.sub(r"[^a-zA-Z0-9_-]", "_", page_url) + ".html"
                    fpath = CACHE_DIR / clean_name

                if not fpath.exists():
                    msg = f"Cached file not found: {fpath} for {page_url}"
                    parser_errors.append(msg)
                    logger.error(msg)
                    continue

                try:
                    html_content = fpath.read_text(encoding="utf-8", errors="replace")
                    lesson_dict = parse_lesson_from_html(
                        html=html_content,
                        source_file=fpath.name,
                        source_url=page_url,
                        page=p_order
                    )
                    parsed_pages_count += 1

                    exercises = lesson_dict.get("pages", [{}])[0].get("exercises", [])
                    for ex in exercises:
                        ex["order"] = exercise_global_order
                        exercise_global_order += 1
                        total_exercises_count += 1

                        for q in ex.get("questions", []):
                            total_questions_count += 1

                            # Radio / checkbox options
                            for opt in q.get("options", []):
                                total_options_count += 1
                                if opt.get("is_correct") is None:
                                    unresolved_answers_count += 1

                            # Gaps
                            for gap in q.get("gaps", []):
                                total_gaps_count += 1
                                if gap.get("correct_answer") is None:
                                    unresolved_answers_count += 1
                                for g_opt in gap.get("options", []):
                                    total_options_count += 1
                                    if g_opt.get("is_correct") is None:
                                        unresolved_answers_count += 1

                    lesson_pages.append({
                        "page": p_order,
                        "source_file": fpath.name,
                        "source_url": page_url,
                        "exercises": exercises
                    })
                except Exception as exc:
                    err_msg = f"Error parsing {fpath.name} ({page_url}): {exc}"
                    parser_errors.append(err_msg)
                    logger.error(err_msg)

            lesson_obj = {
                "lesson_id": lesson_id,
                "title": topic_title,
                "level": level_id,
                "topic": topic_title,
                "status": "parsed",
                "pages": lesson_pages,
                "explanation": {"sections": []},
                "media": [],
                "source": {
                    "provider": "Test-English",
                    "source_url": seed_url,
                    "pages_collected": len(lesson_pages),
                    "source_files": [p["source_file"] for p in lesson_pages]
                }
            }
            lessons_list.append(lesson_obj)

    dataset = {
        "version": "1.0",
        "provider": "Test-English",
        "total_lessons": len(lessons_list),
        "lessons": lessons_list
    }

    # Save JSON preliminary dataset
    OUTPUT_JSON_LOCAL.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_JSON_LOCAL, "w", encoding="utf-8") as f:
        json.dump(dataset, f, ensure_ascii=False, indent=2)
    logger.info(f"Preliminary Universal JSON saved locally: {OUTPUT_JSON_LOCAL}")

    try:
        with open(OUTPUT_JSON_DESKTOP, "w", encoding="utf-8") as f:
            json.dump(dataset, f, ensure_ascii=False, indent=2)
        logger.info(f"Preliminary Universal JSON saved to Desktop: {OUTPUT_JSON_DESKTOP}")
    except Exception as e:
        logger.warning(f"Could not write to Desktop JSON: {e}")

    # Generate Excel CMS workbook
    generate_cms_workbook(dataset, OUTPUT_XLSX_LOCAL)
    try:
        generate_cms_workbook(dataset, OUTPUT_XLSX_DESKTOP)
        logger.info(f"Preliminary Excel CMS workbook saved to Desktop: {OUTPUT_XLSX_DESKTOP}")
    except Exception as e:
        logger.warning(f"Could not write to Desktop Excel: {e}")

    stats = {
        "total_topics_catalog": total_topics_catalog,
        "total_topics_parsed": len(lessons_list),
        "topics_represented_count": len(topics_represented),
        "all_225_topics_represented": (len(lessons_list) == 225 and total_topics_catalog == 225),
        "total_exercises_parsed": total_exercises_count,
        "total_questions": total_questions_count,
        "total_gaps": total_gaps_count,
        "total_options": total_options_count,
        "unresolved_answers_count": unresolved_answers_count,
        "parser_errors_count": len(parser_errors),
        "parser_warnings_count": len(parser_warnings),
        "source_pages_parsed": parsed_pages_count,
        "output_json_local": str(OUTPUT_JSON_LOCAL),
        "output_xlsx_local": str(OUTPUT_XLSX_LOCAL),
        "output_json_desktop": str(OUTPUT_JSON_DESKTOP),
        "output_xlsx_desktop": str(OUTPUT_XLSX_DESKTOP)
    }

    logger.info("=== BATCH OFFLINE PROCESSING SUMMARY ===")
    for k, v in stats.items():
        logger.info(f"  {k}: {v}")

    return stats


if __name__ == "__main__":
    process_entire_corpus()
