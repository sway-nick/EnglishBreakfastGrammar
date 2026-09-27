import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# НАСТРОЙКИ
# ============================================================

INPUT_FILE = Path(r"C:\Users\user\Desktop\universal_lessons.json")
OUTPUT_FILE = Path(r"C:\Users\user\Desktop\english_cms.xlsx")


# ============================================================
# ЗАГРУЗКА JSON
# ============================================================

if not INPUT_FILE.exists():
    print(f"ERROR: File not found:")
    print(INPUT_FILE)
    input("\nPress Enter to exit...")
    raise SystemExit(1)

with open(INPUT_FILE, "r", encoding="utf-8") as f:
    data = json.load(f)


# ============================================================
# СОЗДАНИЕ WORKBOOK
# ============================================================

wb = Workbook()

# Удаляем стандартный лист
default_sheet = wb.active
wb.remove(default_sheet)


# ============================================================
# СТИЛИ
# ============================================================

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

header_alignment = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True
)

body_alignment = Alignment(
    vertical="top",
    wrap_text=True
)


def setup_sheet(ws, headers):
    """
    Создаёт заголовки и базовое оформление листа.
    """

    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws.row_dimensions[1].height = 30


def finish_sheet(ws):
    """
    Настраивает ширину колонок.
    """

    for column_cells in ws.columns:

        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:

            if cell.value is not None:

                value = str(cell.value)

                if len(value) > max_length:
                    max_length = len(value)

        width = min(max(max_length + 2, 12), 50)

        ws.column_dimensions[column_letter].width = width

    for row in ws.iter_rows():

        for cell in row:
            cell.alignment = body_alignment


# ============================================================
# 1. LESSONS
# ============================================================

ws_lessons = wb.create_sheet("Lessons")

lesson_headers = [
    "lesson_id",
    "title",
    "level",
    "topic",
    "status",
    "description",
    "source_provider",
    "source_url",
    "created_at",
    "updated_at"
]

setup_sheet(ws_lessons, lesson_headers)


# ============================================================
# 2. EXERCISES
# ============================================================

ws_exercises = wb.create_sheet("Exercises")

exercise_headers = [
    "exercise_id",
    "lesson_id",
    "page",
    "order",
    "title",
    "instruction",
    "example_source",
    "example_target",
    "status",
    "source_file"
]

setup_sheet(ws_exercises, exercise_headers)


# ============================================================
# 3. QUESTIONS
# ============================================================

ws_questions = wb.create_sheet("Questions")

question_headers = [
    "question_id",
    "exercise_id",
    "lesson_id",
    "order",
    "response_model",
    "content",
    "explanation",
    "difficulty",
    "status"
]

setup_sheet(ws_questions, question_headers)


# ============================================================
# 4. GAPS
# ============================================================

ws_gaps = wb.create_sheet("Gaps")

gap_headers = [
    "gap_id",
    "question_id",
    "gap_order",
    "input_control",
    "correct_answer",
    "accepted_answers",
    "case_sensitive",
    "feedback_correct",
    "feedback_incorrect"
]

setup_sheet(ws_gaps, gap_headers)


# ============================================================
# 5. OPTIONS
# ============================================================

ws_options = wb.create_sheet("Options")

option_headers = [
    "option_id",
    "question_id",
    "gap_id",
    "order",
    "text",
    "value",
    "is_correct"
]

setup_sheet(ws_options, option_headers)


# ============================================================
# 6. EXPLANATIONS
# ============================================================

ws_explanations = wb.create_sheet("Explanations")

explanation_headers = [
    "explanation_id",
    "lesson_id",
    "question_id",
    "section_order",
    "title",
    "content",
    "example",
    "status"
]

setup_sheet(ws_explanations, explanation_headers)


# ============================================================
# ЗАПОЛНЕНИЕ
# ============================================================

lessons = data.get("lessons", [])

lesson_count = 0
exercise_count = 0
question_count = 0
gap_count = 0
option_count = 0
explanation_count = 0


for lesson in lessons:

    lesson_id = lesson.get("lesson_id", "")
    title = lesson.get("title", "")
    level = lesson.get("level", "")
    topic = lesson.get("topic", "")
    status = lesson.get("status", "")

    source = lesson.get("source", {})

    provider = source.get(
        "provider",
        data.get("provider", "")
    )

    source_url = source.get("source_url", "")

    description = lesson.get("description", "")

    ws_lessons.append([
        lesson_id,
        title,
        level,
        topic,
        status,
        description,
        provider,
        source_url,
        "",
        ""
    ])

    lesson_count += 1


    # --------------------------------------------------------
    # PAGES
    # --------------------------------------------------------

    pages = lesson.get("pages", [])

    for page in pages:

        page_number = page.get("page", "")

        exercises = page.get("exercises", [])

        for exercise_index, exercise in enumerate(exercises, start=1):

            exercise_id = exercise.get(
                "exercise_id",
                ""
            )

            instruction = exercise.get(
                "instruction",
                ""
            )

            example = exercise.get(
                "example"
            )

            example_source = ""
            example_target = ""

            if isinstance(example, dict):

                example_source = example.get(
                    "source",
                    ""
                )

                example_target = example.get(
                    "target",
                    ""
                )

            exercise_title = exercise.get(
                "title",
                ""
            )

            exercise_status = exercise.get(
                "status",
                ""
            )

            source_file = exercise.get(
                "source_file",
                ""
            )

            ws_exercises.append([
                exercise_id,
                lesson_id,
                page_number,
                exercise_index,
                exercise_title,
                instruction,
                example_source,
                example_target,
                exercise_status,
                source_file
            ])

            exercise_count += 1


            # ------------------------------------------------
            # QUESTIONS
            # ------------------------------------------------

            questions = exercise.get(
                "questions",
                []
            )

            for question_index, question in enumerate(
                questions,
                start=1
            ):

                question_id = question.get(
                    "question_id",
                    ""
                )

                response_model = question.get(
                    "response_model",
                    ""
                )

                content = question.get(
                    "content",
                    ""
                )

                explanation = question.get(
                    "explanation",
                    ""
                )

                difficulty = question.get(
                    "difficulty",
                    ""
                )

                question_status = question.get(
                    "status",
                    ""
                )

                ws_questions.append([
                    question_id,
                    exercise_id,
                    lesson_id,
                    question_index,
                    response_model,
                    content,
                    explanation,
                    difficulty,
                    question_status
                ])

                question_count += 1


                # ------------------------------------------------
                # QUESTION OPTIONS
                # ------------------------------------------------

                question_options = question.get(
                    "options",
                    []
                )

                for option_index, option in enumerate(
                    question_options,
                    start=1
                ):

                    option_id = option.get(
                        "option_id",
                        ""
                    )

                    option_text = option.get(
                        "text",
                        ""
                    )

                    option_value = option.get(
                        "value",
                        ""
                    )

                    is_correct = option.get(
                        "is_correct"
                    )

                    ws_options.append([
                        option_id,
                        question_id,
                        "",
                        option_index,
                        option_text,
                        option_value,
                        is_correct
                    ])

                    option_count += 1


                # ------------------------------------------------
                # GAPS
                # ------------------------------------------------

                gaps = question.get(
                    "gaps",
                    []
                )

                for gap_index, gap in enumerate(
                    gaps,
                    start=1
                ):

                    gap_id = gap.get(
                        "gap_id",
                        ""
                    )

                    gap_order = gap.get(
                        "order",
                        gap_index
                    )

                    input_control = gap.get(
                        "input_control",
                        ""
                    )

                    correct_answer = gap.get(
                        "correct_answer"
                    )

                    accepted_answers = gap.get(
                        "accepted_answers",
                        []
                    )

                    case_sensitive = gap.get(
                        "case_sensitive",
                        False
                    )

                    feedback_correct = gap.get(
                        "feedback_correct",
                        ""
                    )

                    feedback_incorrect = gap.get(
                        "feedback_incorrect",
                        ""
                    )

                    if isinstance(
                        accepted_answers,
                        list
                    ):
                        accepted_answers_text = " | ".join(
                            str(x)
                            for x in accepted_answers
                        )
                    else:
                        accepted_answers_text = str(
                            accepted_answers
                        )

                    ws_gaps.append([
                        gap_id,
                        question_id,
                        gap_order,
                        input_control,
                        correct_answer,
                        accepted_answers_text,
                        case_sensitive,
                        feedback_correct,
                        feedback_incorrect
                    ])

                    gap_count += 1


                    # --------------------------------------------
                    # GAP OPTIONS
                    # --------------------------------------------

                    gap_options = gap.get(
                        "options",
                        []
                    )

                    for option_index, option in enumerate(
                        gap_options,
                        start=1
                    ):

                        option_id = option.get(
                            "option_id",
                            ""
                        )

                        option_text = option.get(
                            "text",
                            ""
                        )

                        option_value = option.get(
                            "value",
                            ""
                        )

                        is_correct = option.get(
                            "is_correct"
                        )

                        ws_options.append([
                            option_id,
                            question_id,
                            gap_id,
                            option_index,
                            option_text,
                            option_value,
                            is_correct
                        ])

                        option_count += 1


    # --------------------------------------------------------
    # EXPLANATIONS
    # --------------------------------------------------------

    explanation = lesson.get(
        "explanation",
        {}
    )

    sections = explanation.get(
        "sections",
        []
    )

    for section_index, section in enumerate(
        sections,
        start=1
    ):

        explanation_id = section.get(
            "explanation_id",
            f"{lesson_id}_exp_{section_index}"
        )

        section_title = section.get(
            "title",
            ""
        )

        content = section.get(
            "content",
            ""
        )

        example = section.get(
            "example",
            ""
        )

        section_status = section.get(
            "status",
            ""
        )

        ws_explanations.append([
            explanation_id,
            lesson_id,
            "",
            section_index,
            section_title,
            content,
            example,
            section_status
        ])

        explanation_count += 1


# ============================================================
# ФОРМАТИРОВАНИЕ
# ============================================================

for ws in wb.worksheets:
    finish_sheet(ws)


# ============================================================
# СОХРАНЕНИЕ
# ============================================================

wb.save(OUTPUT_FILE)


# ============================================================
# ОТЧЁТ
# ============================================================

print()
print("=" * 70)
print("GOOGLE SHEETS CMS EXPORT")
print("=" * 70)

print()
print(f"Lessons:       {lesson_count}")
print(f"Exercises:     {exercise_count}")
print(f"Questions:     {question_count}")
print(f"Gaps:          {gap_count}")
print(f"Options:       {option_count}")
print(f"Explanations:  {explanation_count}")

print()
print("Sheets created:")
print("  1. Lessons")
print("  2. Exercises")
print("  3. Questions")
print("  4. Gaps")
print("  5. Options")
print("  6. Explanations")

print()
print("FILE:")
print(OUTPUT_FILE)

print()
print("=" * 70)
print("DONE")
print("=" * 70)