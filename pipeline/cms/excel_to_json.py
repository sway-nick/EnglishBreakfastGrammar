#!/usr/bin/env python3
"""
Excel CMS to Universal Lesson JSON Converter.

Part of TASK-004: CMS-to-JSON Export Bridge.
Converts answered or draft Excel CMS workbooks (e.g. english_cms_answered.xlsx)
into canonical Universal Lesson JSON files compatible with the JS Core models,
validation rules, and preview engine.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl


# ---------------------------------------------------------------------------
# HELPERS & NORMALIZATION
# ---------------------------------------------------------------------------

def norm_str(val: Any) -> str:
    """Normalize cell value to stripped string, empty string if None."""
    if val is None:
        return ""
    return str(val).strip()


def parse_int(val: Any, default: int = 1) -> int:
    """Safely parse integer with default fallback."""
    if val is None:
        return default
    try:
        return int(val)
    except (ValueError, TypeError):
        return default


def bool_value(val: Any) -> Optional[bool]:
    """
    Converts Excel cell values to Python bool or None.
    Reuses contract established in cms_validator.py.
    """
    if isinstance(val, bool):
        return val
    if val is None:
        return None
    s = str(val).strip().upper()
    if s == "TRUE":
        return True
    if s == "FALSE":
        return False
    return None


def normalize_placeholders(text: str) -> str:
    """
    Normalizes {{gap_1}} -> {{gap1}} on the export boundary to match
    canonical JS Universal Lesson format.
    """
    if not text:
        return ""
    return re.sub(r"\{\{gap_(\d+)\}\}", r"{{gap\1}}", text)


def reconstruct_accepted_answers(correct_answer: Optional[str], extras_raw: Any) -> List[str]:
    """
    Reconstructs the full list of accepted answers.
    In CMS, 'accepted_answers' contains extras separated by '|'.
    Per Rule & Product Contract, correct_answer must be included first.
    """
    results: List[str] = []
    if correct_answer:
        c = str(correct_answer).strip()
        if c:
            results.append(c)

    if extras_raw is not None:
        extras_str = str(extras_raw).strip()
        if extras_str:
            for part in extras_str.split("|"):
                cleaned = part.strip()
                if cleaned and cleaned not in results:
                    results.append(cleaned)

    return results


# ---------------------------------------------------------------------------
# SHEET LOADER
# ---------------------------------------------------------------------------

def load_sheet(ws) -> List[Dict[str, Any]]:
    """
    Reads rows from openpyxl Worksheet and returns list of dictionaries
    keyed by header names (row 1).
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [norm_str(h) for h in rows[0]]
    items = []

    for row_idx, row in enumerate(rows[1:], start=2):
        # Skip completely empty rows
        if not any(cell is not None and str(cell).strip() != "" for cell in row):
            continue

        item = {}
        for col_idx, header in enumerate(headers):
            if not header:
                continue
            item[header] = row[col_idx] if col_idx < len(row) else None
        item["_row_number"] = row_idx
        items.append(item)

    return items


# ---------------------------------------------------------------------------
# CONVERSION ENGINE
# ---------------------------------------------------------------------------

def convert_workbook_to_lessons(wb: openpyxl.Workbook) -> List[Dict[str, Any]]:
    """
    Converts an openpyxl Workbook into a list of canonical Universal Lesson dicts.
    """
    sheet_names = wb.sheetnames
    required_sheets = ["Lessons", "Exercises", "Questions"]
    for req in required_sheets:
        if req not in sheet_names:
            raise ValueError(f"Missing required sheet '{req}' in workbook. Available: {sheet_names}")

    # Load sheets
    lessons_rows = load_sheet(wb["Lessons"])
    exercises_rows = load_sheet(wb["Exercises"])
    questions_rows = load_sheet(wb["Questions"])
    gaps_rows = load_sheet(wb["Gaps"]) if "Gaps" in sheet_names else []
    options_rows = load_sheet(wb["Options"]) if "Options" in sheet_names else []
    explanations_rows = load_sheet(wb["Explanations"]) if "Explanations" in sheet_names else []

    # 0. Group explanations
    explanations_by_question: Dict[str, str] = {}
    for exp_row in explanations_rows:
        q_id = norm_str(exp_row.get("question_id"))
        content = norm_str(exp_row.get("content")) or norm_str(exp_row.get("title"))
        if q_id and content:
            if q_id in explanations_by_question:
                explanations_by_question[q_id] += "\n\n" + content
            else:
                explanations_by_question[q_id] = content

    # 1. Group options
    options_by_gap: Dict[str, List[Dict[str, Any]]] = {}
    options_by_question: Dict[str, List[Dict[str, Any]]] = {}

    for opt_row in options_rows:
        gap_id = norm_str(opt_row.get("gap_id"))
        q_id = norm_str(opt_row.get("question_id"))

        opt_id = norm_str(opt_row.get("option_id"))
        order = parse_int(opt_row.get("order"), 1)

        val = opt_row.get("value")
        if val is None or str(val).strip() == "":
            val = opt_row.get("text")
        val_str = norm_str(val)

        is_correct = bool_value(opt_row.get("is_correct"))

        opt_obj = {
            "id": opt_id,
            "order": order,
            "value": val_str,
            "is_correct": is_correct,
            "feedback": norm_str(opt_row.get("feedback")),
        }

        if gap_id:
            options_by_gap.setdefault(gap_id, []).append(opt_obj)
        elif q_id:
            options_by_question.setdefault(q_id, []).append(opt_obj)

    # Sort options by order
    for opts in options_by_gap.values():
        opts.sort(key=lambda o: o["order"])
    for opts in options_by_question.values():
        opts.sort(key=lambda o: o["order"])

    # 2. Group gaps
    gaps_by_question: Dict[str, List[Dict[str, Any]]] = {}
    for gap_row in gaps_rows:
        q_id = norm_str(gap_row.get("question_id"))
        gap_id = norm_str(gap_row.get("gap_id"))
        gap_order = parse_int(gap_row.get("gap_order"), 1)

        correct_ans = gap_row.get("correct_answer")
        correct_str = norm_str(correct_ans) if correct_ans is not None else None
        if correct_str == "":
            correct_str = None

        accepted = reconstruct_accepted_answers(correct_str, gap_row.get("accepted_answers"))
        gap_opts = options_by_gap.get(gap_id, [])

        gap_obj = {
            "id": gap_id,
            "order": gap_order,
            "placeholder": f"{{{{gap{gap_order}}}}}",
            "options": gap_opts,
            "correct_answer": correct_str,
            "accepted_answers": accepted,
            "review_required": False,
            "input_control": norm_str(gap_row.get("input_control")),
        }
        if q_id:
            gaps_by_question.setdefault(q_id, []).append(gap_obj)

    for gaps in gaps_by_question.values():
        gaps.sort(key=lambda g: g["order"])

    # 3. Group questions
    questions_by_exercise: Dict[str, List[Dict[str, Any]]] = {}
    for q_row in questions_rows:
        ex_id = norm_str(q_row.get("exercise_id"))
        q_id = norm_str(q_row.get("question_id"))
        q_order = parse_int(q_row.get("order"), 1)
        raw_content = norm_str(q_row.get("content"))
        normalized_text = normalize_placeholders(raw_content)

        resp_model = norm_str(q_row.get("response_model")).lower()
        q_gaps = gaps_by_question.get(q_id, [])
        q_options = options_by_question.get(q_id, [])

        # Determine canonical question type
        if resp_model == "gap" or len(q_gaps) > 0:
            first_control = q_gaps[0].get("input_control", "").lower() if q_gaps else ""
            if first_control == "text":
                q_type = "gap_text"
            else:
                q_type = "gap_select"
        elif resp_model in ("single_choice", "multiple_choice", "true_false", "text_input"):
            q_type = resp_model
        else:
            q_type = resp_model or "gap_select"

        # Clean internal helper fields from gap objects
        cleaned_gaps = []
        for g in q_gaps:
            clean_g = dict(g)
            clean_g.pop("input_control", None)
            cleaned_gaps.append(clean_g)

        exp_text = norm_str(q_row.get("explanation"))
        if not exp_text and q_id in explanations_by_question:
            exp_text = explanations_by_question[q_id]

        q_obj = {
            "id": q_id,
            "order": q_order,
            "type": q_type,
            "text": normalized_text,
            "gaps": cleaned_gaps,
            "options": q_options,
            "hint": norm_str(q_row.get("hint")),
            "feedback": norm_str(q_row.get("feedback")),
            "explanation": exp_text,
            "source": None,
        }
        if ex_id:
            questions_by_exercise.setdefault(ex_id, []).append(q_obj)

    for q_list in questions_by_exercise.values():
        q_list.sort(key=lambda q: q["order"])

    # 4. Group exercises
    exercises_by_lesson: Dict[str, List[Dict[str, Any]]] = {}
    for ex_row in exercises_rows:
        l_id = norm_str(ex_row.get("lesson_id"))
        ex_id = norm_str(ex_row.get("exercise_id"))
        ex_order = parse_int(ex_row.get("order"), 1)

        ex_obj = {
            "id": ex_id,
            "lessonId": l_id,
            "order": ex_order,
            "title": norm_str(ex_row.get("title")),
            "instruction": norm_str(ex_row.get("instruction")),
            "status": norm_str(ex_row.get("status")) or "draft",
            "questions": questions_by_exercise.get(ex_id, []),
            "source": {
                "source_file": norm_str(ex_row.get("source_file")),
                "page": ex_row.get("page"),
            } if (ex_row.get("source_file") or ex_row.get("page")) else None,
        }
        if l_id:
            exercises_by_lesson.setdefault(l_id, []).append(ex_obj)

    for ex_list in exercises_by_lesson.values():
        ex_list.sort(key=lambda e: e["order"])

    # 5. Build lessons
    lessons_list: List[Dict[str, Any]] = []
    for l_idx, l_row in enumerate(lessons_rows, start=1):
        l_id = norm_str(l_row.get("lesson_id"))
        l_order = parse_int(l_row.get("order"), l_idx)

        lesson_obj = {
            "id": l_id,
            "title": norm_str(l_row.get("title")),
            "level": norm_str(l_row.get("level")) or "A1",
            "description": norm_str(l_row.get("description")),
            "status": norm_str(l_row.get("status")) or "draft",
            "order": l_order,
            "version": 1,
            "exercises": exercises_by_lesson.get(l_id, []),
            "source": {
                "provider": norm_str(l_row.get("source_provider")),
                "url": norm_str(l_row.get("source_url")),
            } if (l_row.get("source_provider") or l_row.get("source_url")) else None,
        }
        lessons_list.append(lesson_obj)

    lessons_list.sort(key=lambda l: l["order"])
    return lessons_list


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Convert Excel CMS workbook to canonical Universal Lesson JSON."
    )
    parser.add_argument(
        "--input", "-i",
        required=True,
        help="Path to input .xlsx file (e.g. english_cms_answered.xlsx)",
    )
    parser.add_argument(
        "--output-dir", "-o",
        default="data/json",
        help="Output directory for generated JSON files (default: data/json)",
    )
    parser.add_argument(
        "--split",
        action="store_true",
        help="When set, outputs individual <lesson_id>.json files in addition to universal_lessons.json",
    )

    args = parser.parse_args()

    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8")
            sys.stderr.reconfigure(encoding="utf-8")
        except Exception:
            pass

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"[ERROR] Input file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    print(f"[INFO] Loading workbook: {input_path} ...")
    wb = openpyxl.load_workbook(str(input_path), data_only=True)

    print("[INFO] Converting sheets to Universal Lessons...")
    lessons = convert_workbook_to_lessons(wb)

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1. Write combined universal_lessons.json
    combined_path = out_dir / "universal_lessons.json"
    with open(combined_path, "w", encoding="utf-8") as f:
        json.dump(lessons, f, ensure_ascii=False, indent=2)
    print(f"[SUCCESS] Saved combined lessons ({len(lessons)} lessons): {combined_path}")

    # 2. Optionally write individual lesson files
    if args.split:
        for lesson in lessons:
            l_id = lesson["id"]
            # Sanitize file name
            safe_id = re.sub(r'[^\w\-_\.]', '_', l_id)
            single_path = out_dir / f"{safe_id}.json"
            with open(single_path, "w", encoding="utf-8") as f:
                json.dump(lesson, f, ensure_ascii=False, indent=2)
            print(f"   -> Saved single lesson: {single_path}")

    print("\n[SUCCESS] Conversion completed successfully.")


if __name__ == "__main__":
    main()
