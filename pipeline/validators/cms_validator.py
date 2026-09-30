import sys
from pathlib import Path
from collections import defaultdict
from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


# ============================================================
# CONFIG
# ============================================================

DEFAULT_FILE = r"C:\Users\user\Desktop\english_cms.xlsx"

REQUIRED_SHEETS = [
    "Lessons",
    "Exercises",
    "Questions",
    "Gaps",
    "Options",
    "Explanations",
]


# ============================================================
# HELPERS
# ============================================================

errors = []
warnings = []


def error(message):
    errors.append(message)


def warning(message):
    warnings.append(message)


def norm(value):
    if value is None:
        return ""
    return str(value).strip()


def bool_value(value):
    """
    Converts TRUE/FALSE values from Excel/Google Sheets.
    """
    if isinstance(value, bool):
        return value

    value = norm(value).upper()

    if value == "TRUE":
        return True

    if value == "FALSE":
        return False

    return None


def load_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        return [], {}

    headers = [norm(x) for x in rows[0]]

    data = []

    for row_number, row in enumerate(rows[1:], start=2):
        item = {}

        for i, header in enumerate(headers):
            value = row[i] if i < len(row) else None
            item[header] = value

        item["_excel_row"] = row_number
        data.append(item)

    return data, {h: i for i, h in enumerate(headers)}


def require_columns(sheet_name, data, required):
    if not data:
        warning(f"{sheet_name}: sheet is empty")
        return

    existing = set(data[0].keys())

    for column in required:
        if column not in existing:
            error(
                f"{sheet_name}: missing required column '{column}'"
            )


def check_unique(data, column, sheet_name):
    seen = {}

    for row in data:
        value = norm(row.get(column))

        if not value:
            continue

        if value in seen:
            error(
                f"{sheet_name}: duplicate {column}='{value}' "
                f"at rows {seen[value]} and {row['_excel_row']}"
            )
        else:
            seen[value] = row["_excel_row"]


# ============================================================
# MAIN
# ============================================================

def validate_excel(file_path):

    print("=" * 70)
    print("ENGLISH CMS — FINAL VALIDATOR")
    print("=" * 70)
    print()

    if not Path(file_path).exists():
        print(f"ERROR: file not found:")
        print(file_path)
        return 1

    print(f"FILE:")
    print(file_path)
    print()

    try:
        wb = load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"ERROR: cannot open Excel file:")
        print(e)
        return 1

    # --------------------------------------------------------
    # SHEETS
    # --------------------------------------------------------

    print("CHECK 1 — SHEETS")
    print("-" * 70)

    for sheet in REQUIRED_SHEETS:
        if sheet not in wb.sheetnames:
            error(f"Missing sheet: {sheet}")
        else:
            print(f"  OK  {sheet}")

    print()

    if errors:
        return finish()

    # --------------------------------------------------------
    # LOAD DATA
    # --------------------------------------------------------

    lessons, _ = load_sheet(wb["Lessons"])
    exercises, _ = load_sheet(wb["Exercises"])
    questions, _ = load_sheet(wb["Questions"])
    gaps, _ = load_sheet(wb["Gaps"])
    options, _ = load_sheet(wb["Options"])
    explanations, _ = load_sheet(wb["Explanations"])

    # --------------------------------------------------------
    # REQUIRED COLUMNS
    # --------------------------------------------------------

    print("CHECK 2 — COLUMNS")
    print("-" * 70)

    require_columns(
        "Lessons",
        lessons,
        [
            "lesson_id",
            "title",
            "level",
            "topic",
            "status",
            "source_provider",
            "source_url",
        ],
    )

    require_columns(
        "Exercises",
        exercises,
        [
            "exercise_id",
            "lesson_id",
            "page",
            "order",
            "title",
            "instruction",
            "example_source",
            "example_target",
            "status",
        ],
    )

    require_columns(
        "Questions",
        questions,
        [
            "question_id",
            "exercise_id",
            "lesson_id",
            "order",
            "response_model",
            "content",
            "status",
        ],
    )

    require_columns(
        "Gaps",
        gaps,
        [
            "gap_id",
            "question_id",
            "gap_order",
            "input_control",
            "correct_answer",
            "accepted_answers",
            "case_sensitive",
            "feedback_correct",
            "feedback_incorrect",
        ],
    )

    require_columns(
        "Options",
        options,
        [
            "option_id",
            "question_id",
            "gap_id",
            "order",
            "text",
            "value",
            "is_correct",
        ],
    )

    require_columns(
        "Explanations",
        explanations,
        [
            "explanation_id",
            "lesson_id",
            "question_id",
            "section_order",
            "title",
            "content",
            "example",
            "status",
        ],
    )

    if errors:
        return finish()

    print("  Required columns: OK")
    print()

    # --------------------------------------------------------
    # UNIQUE IDS
    # --------------------------------------------------------

    print("CHECK 3 — UNIQUE IDs")
    print("-" * 70)

    check_unique(lessons, "lesson_id", "Lessons")
    check_unique(exercises, "exercise_id", "Exercises")
    check_unique(questions, "question_id", "Questions")
    check_unique(gaps, "gap_id", "Gaps")
    check_unique(options, "option_id", "Options")
    check_unique(explanations, "explanation_id", "Explanations")

    if not errors:
        print("  IDs: OK")

    print()

    # --------------------------------------------------------
    # CREATE LOOKUPS
    # --------------------------------------------------------

    lesson_map = {
        norm(x["lesson_id"]): x
        for x in lessons
        if norm(x.get("lesson_id"))
    }

    exercise_map = {
        norm(x["exercise_id"]): x
        for x in exercises
        if norm(x.get("exercise_id"))
    }

    question_map = {
        norm(x["question_id"]): x
        for x in questions
        if norm(x.get("question_id"))
    }

    gap_map = {
        norm(x["gap_id"]): x
        for x in gaps
        if norm(x.get("gap_id"))
    }

    # --------------------------------------------------------
    # LESSONS
    # --------------------------------------------------------

    print("CHECK 4 — LESSONS")
    print("-" * 70)

    for row in lessons:

        lesson_id = norm(row.get("lesson_id"))

        if not lesson_id:
            error(
                f"Lessons row {row['_excel_row']}: empty lesson_id"
            )

        if not norm(row.get("title")):
            error(
                f"Lessons row {row['_excel_row']}: empty title"
            )

    print(f"  Lessons: {len(lessons)}")
    print()

    # --------------------------------------------------------
    # EXERCISES
    # --------------------------------------------------------

    print("CHECK 5 — EXERCISES")
    print("-" * 70)

    for row in exercises:

        exercise_id = norm(row.get("exercise_id"))
        lesson_id = norm(row.get("lesson_id"))

        if not exercise_id:
            error(
                f"Exercises row {row['_excel_row']}: empty exercise_id"
            )

        if lesson_id not in lesson_map:
            error(
                f"Exercises row {row['_excel_row']}: "
                f"unknown lesson_id '{lesson_id}'"
            )

        if not norm(row.get("instruction")):
            warning(
                f"Exercises row {row['_excel_row']}: "
                f"empty instruction"
            )

    print(f"  Exercises: {len(exercises)}")
    print()

    # --------------------------------------------------------
    # QUESTIONS
    # --------------------------------------------------------

    print("CHECK 6 — QUESTIONS")
    print("-" * 70)

    valid_models = {
        "gap",
        "single_choice",
        "multiple_choice",
    }

    for row in questions:

        question_id = norm(row.get("question_id"))
        exercise_id = norm(row.get("exercise_id"))
        lesson_id = norm(row.get("lesson_id"))
        response_model = norm(row.get("response_model"))

        if not question_id:
            error(
                f"Questions row {row['_excel_row']}: "
                f"empty question_id"
            )

        if exercise_id not in exercise_map:
            error(
                f"Question {question_id}: "
                f"unknown exercise_id '{exercise_id}'"
            )

        if lesson_id not in lesson_map:
            error(
                f"Question {question_id}: "
                f"unknown lesson_id '{lesson_id}'"
            )

        if response_model not in valid_models:
            error(
                f"Question {question_id}: "
                f"invalid response_model '{response_model}'"
            )

        if not norm(row.get("content")):
            error(
                f"Question {question_id}: empty content"
            )

    print(f"  Questions: {len(questions)}")
    print()

    # --------------------------------------------------------
    # GAPS
    # --------------------------------------------------------

    print("CHECK 7 — GAPS")
    print("-" * 70)

    gaps_by_question = defaultdict(list)

    for row in gaps:

        gap_id = norm(row.get("gap_id"))
        question_id = norm(row.get("question_id"))

        if not gap_id:
            error(
                f"Gaps row {row['_excel_row']}: empty gap_id"
            )

        if question_id not in question_map:
            error(
                f"Gap {gap_id}: unknown question_id "
                f"'{question_id}'"
            )

        gaps_by_question[question_id].append(row)

        correct_answer = norm(row.get("correct_answer"))

        if not correct_answer:
            error(
                f"Gap {gap_id}: empty correct_answer"
            )

        accepted_answers = norm(row.get("accepted_answers"))

        if not accepted_answers:
            error(
                f"Gap {gap_id}: empty accepted_answers"
            )

        input_control = norm(row.get("input_control"))

        if input_control not in {"select", "text"}:
            error(
                f"Gap {gap_id}: invalid input_control "
                f"'{input_control}'"
            )

    print(f"  Gaps: {len(gaps)}")
    print()

    # --------------------------------------------------------
    # GAP PLACEHOLDERS
    # --------------------------------------------------------

    print("CHECK 8 — GAP PLACEHOLDERS")
    print("-" * 70)

    import re

    for question_id, question in question_map.items():

        content = norm(question.get("content"))

        placeholders = re.findall(
            r"\{\{gap_(\d+)\}\}",
            content
        )

        expected = len(placeholders)
        actual = len(gaps_by_question.get(question_id, []))

        if expected != actual:

            error(
                f"Question {question_id}: "
                f"placeholders={expected}, gaps={actual}"
            )

    if not errors:
        print("  Gap placeholders: OK")

    print()

    # --------------------------------------------------------
    # OPTIONS
    # --------------------------------------------------------

    print("CHECK 9 — OPTIONS")
    print("-" * 70)

    options_by_question = defaultdict(list)
    options_by_gap = defaultdict(list)

    for row in options:

        option_id = norm(row.get("option_id"))
        question_id = norm(row.get("question_id"))
        gap_id = norm(row.get("gap_id"))

        if not option_id:
            error(
                f"Options row {row['_excel_row']}: "
                f"empty option_id"
            )

        if question_id not in question_map:
            error(
                f"Option {option_id}: "
                f"unknown question_id '{question_id}'"
            )

        if not norm(row.get("text")):
            error(
                f"Option {option_id}: empty text"
            )

        correct = bool_value(row.get("is_correct"))

        if correct is None:
            error(
                f"Option {option_id}: "
                f"is_correct must be TRUE/FALSE"
            )

        options_by_question[question_id].append(row)

        if gap_id:
            if gap_id not in gap_map:
                error(
                    f"Option {option_id}: "
                    f"unknown gap_id '{gap_id}'"
                )

            options_by_gap[gap_id].append(row)

    print(f"  Options: {len(options)}")
    print()

    # --------------------------------------------------------
    # OPTION CORRECTNESS
    # --------------------------------------------------------

    print("CHECK 10 — CORRECT OPTIONS")
    print("-" * 70)

    for question_id, question in question_map.items():

        model = norm(question.get("response_model"))

        q_options = options_by_question.get(question_id, [])

        if model == "single_choice":

            # single_choice MUST NOT have gap_id
            for option in q_options:

                gap_id = norm(option.get("gap_id"))

                if gap_id:
                    error(
                        f"Question {question_id}: "
                        f"single_choice option "
                        f"{option['option_id']} has gap_id"
                    )

            true_count = sum(
                bool_value(x.get("is_correct")) is True
                for x in q_options
            )

            if true_count != 1:
                error(
                    f"Question {question_id}: "
                    f"single_choice has {true_count} "
                    f"correct options; expected 1"
                )

        elif model == "multiple_choice":

            # multiple_choice MUST NOT have gap_id
            for option in q_options:

                gap_id = norm(option.get("gap_id"))

                if gap_id:
                    error(
                        f"Question {question_id}: "
                        f"multiple_choice option "
                        f"{option['option_id']} has gap_id"
                    )

            true_count = sum(
                bool_value(x.get("is_correct")) is True
                for x in q_options
            )

            if true_count < 1:
                error(
                    f"Question {question_id}: "
                    f"multiple_choice has {true_count} "
                    f"correct options; expected >= 1"
                )

        elif model == "gap":

            q_gaps = gaps_by_question.get(question_id, [])

            for gap in q_gaps:

                gap_id = norm(gap.get("gap_id"))
                input_control = norm(gap.get("input_control"))

                gap_options = options_by_gap.get(gap_id, [])

                if input_control == "select":

                    if not gap_options:
                        error(
                            f"Gap {gap_id}: "
                            f"select gap has no options"
                        )

                    true_count = sum(
                        bool_value(x.get("is_correct")) is True
                        for x in gap_options
                    )

                    if true_count != 1:
                        error(
                            f"Gap {gap_id}: "
                            f"has {true_count} correct options; "
                            f"expected 1"
                        )

                elif input_control == "text":

                    if gap_options:
                        error(
                            f"Gap {gap_id}: "
                            f"text gap must not have options"
                        )

    if not errors:
        print("  Correct option logic: OK")

    print()

    # --------------------------------------------------------
    # GAPS <-> OPTIONS
    # --------------------------------------------------------

    print("CHECK 11 -- GAPS <-> OPTIONS")
    print("-" * 70)

    for gap_id, gap in gap_map.items():

        input_control = norm(gap.get("input_control"))

        correct_answer = norm(
            gap.get("correct_answer")
        ).lower()

        if input_control != "select":
            continue

        gap_options = options_by_gap.get(gap_id, [])

        correct_options = [
            x for x in gap_options
            if bool_value(x.get("is_correct")) is True
        ]

        if len(correct_options) != 1:
            continue

        option_text = norm(
            correct_options[0].get("text")
        ).lower()
        option_val = norm(
            correct_options[0].get("value")
        ).lower()

        if option_text != correct_answer and option_val != correct_answer:

            error(
                f"Gap {gap_id}: "
                f"Gaps.correct_answer='{gap.get('correct_answer')}' "
                f"but correct option text='{correct_options[0].get('text')}'"
            )

    if not errors:
        print("  Gaps ↔ Options: OK")

    print()

    # --------------------------------------------------------
    # ORDER CHECK
    # --------------------------------------------------------

    print("CHECK 12 — ORDERS")
    print("-" * 70)

    for question_id, q_gaps in gaps_by_question.items():

        orders = []

        for gap in q_gaps:
            try:
                orders.append(int(gap["gap_order"]))
            except:
                error(
                    f"Gap {gap.get('gap_id')}: "
                    f"invalid gap_order"
                )

        if orders:

            expected = list(range(1, len(orders) + 1))

            if sorted(orders) != expected:
                error(
                    f"Question {question_id}: "
                    f"gap_order is {orders}, "
                    f"expected {expected}"
                )

    print("  Order checks completed")
    print()

    # --------------------------------------------------------
    # STATISTICS
    # --------------------------------------------------------

    print("=" * 70)
    print("STATISTICS")
    print("=" * 70)

    print(f"Lessons:       {len(lessons)}")
    print(f"Exercises:     {len(exercises)}")
    print(f"Questions:     {len(questions)}")
    print(f"Gaps:          {len(gaps)}")
    print(f"Options:       {len(options)}")
    print(f"Explanations:  {len(explanations)}")

    print()

    return finish()


# ============================================================
# FINISH
# ============================================================

def finish():

    print("=" * 70)
    print("VALIDATION RESULT")
    print("=" * 70)

    print(f"Errors:   {len(errors)}")
    print(f"Warnings: {len(warnings)}")
    print()

    if errors:

        print("ERRORS")
        print("-" * 70)

        for i, message in enumerate(errors, 1):
            print(f"{i}. {message}")

        print()
        print("VALIDATION FAILED")
        return 1

    if warnings:

        print("WARNINGS")
        print("-" * 70)

        for i, message in enumerate(warnings, 1):
            print(f"{i}. {message}")

        print()

    print("VALIDATION PASSED")
    return 0


# ============================================================
# CLI
# ============================================================

if __name__ == "__main__":

    file_path = (
        sys.argv[1]
        if len(sys.argv) > 1
        else DEFAULT_FILE
    )

    sys.exit(
        validate_excel(file_path)
    )